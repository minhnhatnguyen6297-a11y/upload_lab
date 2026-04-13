from __future__ import annotations

import tempfile
import unicodedata
import unittest
from pathlib import Path

from docx import Document

from extract_contract import extract, find_tai_san, guess_loai_tai_san


def make_docx(path: Path, *paragraphs: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(str(path))
    return path


def make_docx_with_blocks(path: Path, blocks: list[tuple[str, object]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for block_type, payload in blocks:
        if block_type == "p":
            doc.add_paragraph(str(payload))
            continue
        if block_type == "table":
            rows = list(payload)
            col_count = max(len(row) for row in rows)
            table = doc.add_table(rows=len(rows), cols=col_count)
            for row_idx, row in enumerate(rows):
                for col_idx, value in enumerate(row):
                    table.cell(row_idx, col_idx).text = str(value)
            continue
        raise ValueError(f"Unsupported block type: {block_type}")
    doc.save(str(path))
    return path


class UploadLabExtractContractTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.root = Path(self.tempdir.name)

    def test_transfer_title_ignores_body_huy_bo_clause(self):
        docx_path = make_docx(
            self.root / "transfer.docx",
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT",
            "Chúng tôi gồm có:",
            "I. BÊN CHUYỂN NHƯỢNG: (Bên A)",
            "Ông: Nguyễn Văn A Sinh ngày: 01/01/1980",
            "Căn cước công dân số: 012345678901 do Bộ Công an cấp ngày 01/01/2024;",
            "Thường trú tại: thôn A, xã B.",
            "II. BÊN NHẬN CHUYỂN NHƯỢNG: (Bên B)",
            "Ông: Nguyễn Văn B Sinh ngày: 02/02/1981",
            "Căn cước công dân số: 012345678902 do Bộ Công an cấp ngày 02/02/2024;",
            "Thường trú tại: thôn C, xã D.",
            "ĐIỀU 1: ĐỐI TƯỢNG CỦA HỢP ĐỒNG",
            "Đối tượng của Hợp đồng này là toàn bộ quyền sử dụng đất của bên A có địa chỉ tại: thôn A, xã B.",
            "- Thửa đất số: 10",
            "ĐIỀU 8: ĐIỀU KHOẢN CUỐI CÙNG",
            "Việc sửa đổi, bổ sung hoặc hủy bỏ hợp đồng này chỉ có giá trị khi được hai bên lập thành văn bản.",
            "Số công chứng 428/2026/CCGD",
        )

        payload = extract(docx_path)

        self.assertEqual(payload["raw"]["document_kind"], "transfer_contract")
        self.assertEqual(payload["web_form"]["ten_hop_dong"], "Hợp đồng chuyển nhượng quyền sử dụng đất")

    def test_find_tai_san_commitment_stops_before_cam_ket_xac_nhan(self):
        text = "\n".join(
            [
                "VĂN BẢN CAM KẾT TÀI SẢN RIÊNG",
                "Hiện nay, ông Nguyễn Văn Nam đang làm các thủ tục để nhận chuyển nhượng quyền sử dụng đất có địa chỉ tại: thôn A, xã B.",
                "- Thửa đất số: 20",
                "- Diện tích: 200 m2",
                "- Giấy chứng nhận quyền sử dụng đất quyền sở hữu nhà ở và tài sản khác gắn liền với đất số: AB 123456.",
                "Hai vợ chồng chúng tôi cam kết và xác nhận: Toàn bộ số tiền này là tài sản riêng.",
            ]
        )

        tai_san = find_tai_san(unicodedata.normalize("NFD", text))

        self.assertIn("Thửa đất số: 20", tai_san)
        self.assertNotIn("Hai vợ chồng chúng tôi cam kết", tai_san)

    def test_extract_commitment_keeps_both_spouses_and_classifies_land_only(self):
        docx_path = make_docx(
            self.root / "cam_ket_tai_san_rieng.docx",
            "VĂN BẢN CAM KẾT TÀI SẢN RIÊNG",
            "Chúng tôi gồm có:",
            "- Người chồng – Ông: Nguyễn Văn Nam Sinh ngày: 02/02/1982",
            "Căn cước công dân số: 036082000989 do Bộ Công an cấp ngày 02/07/2021",
            "Thường trú tại: Thôn Hoàng Mẫu, xã Vạn Thắng, tỉnh Ninh Bình;",
            "- Người vợ - Bà: Nguyễn Thị Oanh Sinh ngày: 09/05/1985",
            "Căn cước số: 036185021354 do Bộ Công an cấp ngày 12/02/2025;",
            "Nơi cư trú tại: Thôn Hoàng Mẫu, xã Vạn Thắng, tỉnh Ninh Bình.",
            "Chúng tôi là vợ chồng theo Giấy chứng nhận kết hôn số 17.",
            "Nay chúng tôi lập văn bản này để cam kết và chịu trách nhiệm trước pháp luật về những nội dung sau đây:",
            "Hiện nay, ông Nguyễn Văn Nam đang làm các thủ tục để nhận chuyển nhượng quyền sử dụng đất có địa chỉ tại: thôn Đại Lộc Trung, xã Yên Chính, huyện Ý Yên, tỉnh Nam Định Giấy chứng nhận quyền sử dụng đất quyền sở hữu nhà ở và tài sản khác gắn liền với đất số: CY 921651.",
            "- Thửa đất số: 531; Tờ bản đồ số: 22;",
            "- Diện tích: 660,0 m2;",
            "Hai vợ chồng chúng tôi cam kết và xác nhận: Toàn bộ số tiền này là tài sản riêng của ông Nguyễn Văn Nam.",
            "Số công chứng 161/2026/CCGD",
        )

        payload = extract(docx_path)
        duong_su = payload["web_form"]["duong_su"]
        tai_san = payload["web_form"]["tai_san"]

        self.assertEqual(payload["raw"]["document_kind"], "asset_commitment")
        self.assertEqual(payload["web_form"]["ten_hop_dong"], "Văn bản cam kết tài sản riêng")
        self.assertEqual(payload["web_form"]["nhom_hop_dong"], "Thoả thuận - Cam kết")
        self.assertEqual(payload["web_form"]["loai_tai_san"], "Đất đai không có tài sản")
        self.assertIn("Nguyễn Văn Nam", duong_su)
        self.assertIn("Nguyễn Thị Oanh", duong_su)
        self.assertNotIn("Hai vợ chồng chúng tôi cam kết", tai_san)

    def test_guess_loai_tai_san_detects_real_attached_asset_details_after_stripping_certificate_heading_noise(self):
        tai_san = "\n".join(
            [
                "Quyền sử dụng đất có địa chỉ tại: thôn A, xã B.",
                "Giấy chứng nhận quyền sử dụng đất, quyền sở hữu nhà ở và tài sản khác gắn liền với đất số: AB 123456.",
                "Nhà ở:",
                "- Diện tích xây dựng: 120 m2",
                "- Diện tích sàn: 240 m2",
                "- Kết cấu: Tường gạch, mái bê tông cốt thép",
            ]
        )

        loai_tai_san = guess_loai_tai_san(tai_san, "Văn bản cam kết tài sản riêng")

        self.assertEqual(loai_tai_san, "Đất đai có tài sản")

    def test_generic_title_strips_duoc_giao_ket_boi_suffix(self):
        docx_path = make_docx(
            self.root / "uy_quyen.docx",
            "CHỨNG NHẬN:",
            "Hợp đồng ủy quyền này được giao kết bởi:",
            "I. BÊN ỦY QUYỀN",
            "Bà Tạ Thị Chứ; Sinh ngày: 08/01/1956;",
            "Căn cước công dân số 036156012959 do Bộ Công an cấp ngày 10/01/2023;",
            "Thường trú tại: Thôn Tâm Bình, xã Yên Cường, tỉnh Ninh Bình.",
            "II. BÊN ĐƯỢC ỦY QUYỀN",
            "Ông Nguyễn Đức Hiệp; Sinh ngày: 17/05/1982;",
            "Căn cước công dân số 036082016119 do Bộ Công an cấp ngày 09/08/2021;",
            "Thường trú tại: Thôn Tâm Bình, xã Yên Cường, tỉnh Ninh Bình.",
            "Số công chứng 372/2026/CCGD",
        )

        payload = extract(docx_path)

        self.assertEqual(payload["web_form"]["ten_hop_dong"], "Hợp đồng ủy quyền")

    def test_party_blocks_keep_shared_tru_tai_lines_for_each_person(self):
        docx_path = make_docx(
            self.root / "shared_address.docx",
            "HỢP ĐỒNG CHUYỂN NHƯỢNG QUYỀN SỬ DỤNG ĐẤT VÀ TÀI SẢN GẮN LIỀN VỚI ĐẤT",
            "Chúng tôi gồm có:",
            "I. BÊN CHUYỂN NHƯỢNG: (Bên A)",
            "1. Ông Nguyễn Huy Cận Sinh ngày: 17/12/1979",
            "Căn cước công dân số: 037079010407 do Bộ Công an cấp ngày 19/02/2025;",
            "2. Bà Nguyễn Thị Thu Thủy Sinh ngày: 27/06/1984",
            "Căn cước công dân số: 022184004248 do Bộ Công an cấp ngày 21/10/2025;",
            "Cùng cư trú tại: thôn A, xã B.",
            "II. BÊN NHẬN CHUYỂN NHƯỢNG: (Bên B)",
            "1. Ông Đoàn Văn Khánh Sinh ngày: 07/04/1987",
            "Căn cước công dân số: 036087026862 do Bộ Công an cấp ngày 11/08/2021;",
            "2. Bà Nguyễn Thị Đức Hạnh Sinh ngày: 14/12/1989",
            "Căn cước công dân số: 036189016258 do Bộ Công an cấp ngày 02/05/2022;",
            "Cùng cư trú tại: thôn C, xã D.",
            "ĐIỀU 1: ĐỐI TƯỢNG CỦA HỢP ĐỒNG",
            "Đối tượng của Hợp đồng này là toàn bộ quyền sử dụng đất của bên A có địa chỉ tại: thôn A, xã B.",
            "- Thửa đất số: 99",
            "Số công chứng 500/2026/CCGD",
        )

        payload = extract(docx_path)
        duong_su = payload["web_form"]["duong_su"]
        nguoi_yeu_cau = payload["web_form"]["nguoi_yeu_cau"]

        self.assertEqual(payload["web_form"]["ten_hop_dong"], "Hợp đồng chuyển nhượng quyền sử dụng đất")
        self.assertEqual(duong_su.count("Cùng cư trú tại: thôn A, xã B."), 2)
        self.assertEqual(duong_su.count("Cùng cư trú tại: thôn C, xã D."), 2)
        self.assertIn("Cùng cư trú tại: thôn C, xã D.", nguoi_yeu_cau)


    def test_extract_mortgage_summary_keeps_table_order_and_uses_mortgage_party_labels(self):
        docx_path = make_docx_with_blocks(
            self.root / "the_chap_tom_tat.docx",
            [
                ("p", "LỜI CHỨNG CỦA CÔNG CHỨNG VIÊN"),
                ("p", "Hợp đồng thế chấp quyền sử dụng đất được giao kết giữa:"),
                ("p", "I. Bên A: BÊN THẾ CHẤP"),
                (
                    "table",
                    [
                        ("1. Ông: Vũ Duy Quang", "Ngày sinh: 15/04/1981"),
                        ("Căn cước công dân 036081005455 do Bộ Công an cấp ngày 25/04/2021",),
                        ("Địa chỉ nơi cư trú: xã Ý Yên, tỉnh Ninh Bình",),
                        ("2. và vợ là Bà: Phạm Thị Lan", "Ngày sinh: 06/05/1983"),
                        ("Căn cước công dân: 036183009426 do Bộ Công an cấp ngày 06/08/2023",),
                        ("Địa chỉ nơi cư trú: xã Ý Yên, tỉnh Ninh Bình",),
                    ],
                ),
                ("p", "II. Bên B: BÊN NHẬN THẾ CHẤP"),
                ("p", "NGÂN HÀNG THƯƠNG MẠI CỔ PHẦN NGOẠI THƯƠNG VIỆT NAM – CHI NHÁNH NAM ĐỊNH"),
                ("p", "Địa chỉ đăng ký: Số 629 Trần Hưng Đạo, phường Nam Định, tỉnh Ninh Bình."),
                ("p", "Đại diện: Ông Nguyễn Hữu Dụng Chức vụ: Phó giám đốc chi nhánh"),
                ("p", "Căn cước công dân số 036081017402 do Bộ Công an cấp ngày 02/07/2021"),
                ("p", "Số công chứng 473/2026/CCGD"),
            ],
        )

        payload = extract(docx_path)
        duong_su = payload["web_form"]["duong_su"]

        self.assertEqual(payload["raw"]["document_kind"], "mortgage_contract")
        self.assertEqual(payload["web_form"]["ten_hop_dong"], "Hợp đồng thế chấp quyền sử dụng đất")
        self.assertEqual(payload["web_form"]["nhom_hop_dong"], "Cầm cố - Thế chấp - Vay")
        self.assertEqual(payload["web_form"]["tai_san"], "Quyền sử dụng đất")
        self.assertEqual(payload["web_form"]["loai_tai_san"], "Đất đai không có tài sản")
        self.assertIn("BÊN THẾ CHẤP:", duong_su)
        self.assertIn("BÊN NHẬN THẾ CHẤP:", duong_su)
        self.assertIn("NGÂN HÀNG THƯƠNG MẠI CỔ PHẦN NGOẠI THƯƠNG VIỆT NAM", duong_su)
        self.assertIn("Nguyễn Hữu Dụng", payload["web_form"]["nguoi_yeu_cau"])
        self.assertEqual(len(payload["raw"]["ben_a"]["nguoi"]), 2)
        self.assertEqual(len(payload["raw"]["ben_b"]["nguoi"]), 1)

    def test_extract_inheritance_partition_uses_single_heir_group_and_stops_asset_before_heirs_section(self):
        docx_path = make_docx(
            self.root / "phan_chia_di_san.docx",
            "VĂN BẢN PHÂN CHIA DI SẢN",
            "Chúng tôi là những người được hưởng di sản theo pháp luật của ông Nguyễn Văn Hỷ:",
            "Bà Nguyễn Thị Tuyết Sinh ngày: 01/01/1961",
            "Căn cước công dân số 036161014031 do Bộ Công an cấp ngày 28/06/2021;",
            "Thường trú tại: 02 Bà Huyện Thanh Quan, phường Lâm Viên – Đà Lạt, tỉnh Lâm Đồng.",
            "Bà Nguyễn Thị Nga Sinh ngày: 02/04/1964",
            "Căn cước công dân số 036164017015 do Bộ Công an cấp ngày 14/03/2024;",
            "Thường trú tại: Số 06/Đ2 Phan Chu Trinh, phường Lâm Viên – Đà Lạt, tỉnh Lâm Đồng.",
            "Bà Nguyễn Thị Thu Là Sinh ngày: 04/10/1972",
            "Căn cước công dân số 036172012182 do Bộ Công an cấp ngày 28/06/2021;",
            "Thường trú tại: 71C/2 Bùi Thị Xuân, phường Lâm Viên – Đà Lạt, tỉnh Lâm Đồng.",
            "Chúng tôi tự nguyện lập Văn bản này với nội dung như sau:",
            "Người để lại di sản:",
            "Ông Nguyễn Văn Hỷ; Sinh năm: 1928; chết ngày 09/10/2019.",
            "Di sản:",
            "Di sản của ông Nguyễn Văn Hỷ để lại là:",
            "Quyền sử dụng đất của ông Nguyễn Văn Hỷ có địa chỉ tại: xã Ý Yên, tỉnh Ninh Bình theo Giấy chứng nhận quyền sử dụng đất số: A 692002.",
            "- Thửa đất số: 65; Tờ bản đồ số: 21",
            "- Diện tích: 190 m²",
            "- Mục đích sử dụng: Đất ở",
            "Người thừa kế:",
            "Những người thừa kế của ông Nguyễn Văn Hỷ gồm vợ và các con.",
            "Nội dung phân chia di sản.",
            "Bằng Văn bản này, chúng tôi xin nhận kỷ phần thừa kế.",
            "LỜI CHỨNG CỦA CÔNG CHỨNG VIÊN",
            "Số công chứng 2433.2025/PCDS/CCGD.",
        )

        payload = extract(docx_path)
        duong_su = payload["web_form"]["duong_su"]
        tai_san = payload["web_form"]["tai_san"]

        self.assertEqual(payload["raw"]["document_kind"], "inheritance_partition")
        self.assertEqual(payload["web_form"]["ten_hop_dong"], "Văn bản phân chia di sản")
        self.assertEqual(payload["web_form"]["nhom_hop_dong"], "Thừa kế (khai nhận - phân chia di sản thừa kế )")
        self.assertEqual(payload["web_form"]["so_cong_chung"], "2433/2025")
        self.assertTrue(duong_su.startswith("NHỮNG NGƯỜI HƯỞNG DI SẢN:"))
        self.assertIn("Nguyễn Thị Tuyết", duong_su)
        self.assertIn("Nguyễn Thị Nga", duong_su)
        self.assertIn("Nguyễn Thị Thu Là", duong_su)
        self.assertNotIn("BÊN B:", duong_su)
        self.assertNotIn("Người để lại di sản", duong_su)
        self.assertIn("Quyền sử dụng đất của ông Nguyễn Văn Hỷ có địa chỉ tại", tai_san)
        self.assertNotIn("Người thừa kế:", tai_san)
        self.assertNotIn("Nội dung phân chia di sản", tai_san)

    def test_extract_inheritance_refusal_keeps_refuser_and_multiple_assets_without_oath_section(self):
        docx_path = make_docx_with_blocks(
            self.root / "tu_choi_di_san.docx",
            [
                ("p", "VĂN BẢN TỪ CHỐI NHẬN DI SẢN"),
                ("p", "Tôi là: Bà Hoàng Thị Hợi; Sinh ngày: 02/02/1945;"),
                ("p", "Căn cước công dân số 037145004692 do Bộ Công an cấp ngày 02/07/2021;"),
                ("p", "Thường trú tại: Thôn Hưng Thượng, xã Ý Yên, tỉnh Ninh Bình."),
                ("p", "Nay, tôi tự nguyện lập Văn bản này với nội dung như sau:"),
                ("p", "Chồng tôi là ông Nguyễn Văn Hỷ; Sinh năm: 1928; chết ngày 09/10/2019."),
                ("p", "Theo quy định của pháp luật, tôi là người được hưởng thừa kế đối với di sản mà ông Nguyễn Văn Hỷ để lại, đó là:"),
                ("p", "1. Tài sản thứ nhất:"),
                ("p", "Quyền sử dụng đất của ông Nguyễn Văn Hỷ có địa chỉ tại: xã Ý Yên, tỉnh Ninh Bình theo Giấy chứng nhận quyền sử dụng đất số: A 692002."),
                ("p", "- Thửa đất số: 65; Tờ bản đồ số: 21"),
                ("p", "- Diện tích: 190 m²"),
                ("p", "2. Tài sản thứ hai:"),
                ("p", "Phần quyền sử dụng đất của ông Nguyễn Văn Hỷ trong khối tài sản chung với vợ là bà Hoàng Thị Hợi có địa chỉ tại: xã Ý Yên, tỉnh Ninh Bình theo Giấy chứng nhận quyền sử dụng đất, quyền sở hữu nhà ở và tài sản khác gắn liền với đất số: CK 902585."),
                (
                    "table",
                    [
                        ("Tờ bản đồ số", "Thửa đất số", "Diện tích (m2)", "Mục đích sử dụng"),
                        ("19", "122(12)", "216.0", "Đất trồng lúa"),
                        ("46", "4(8)", "36.0", "Đất trồng lúa"),
                    ],
                ),
                ("p", "Bằng Văn bản này, tôi – Hoàng Thị Hợi tự nguyện từ chối nhận kỷ phần thừa kế mà mình được hưởng."),
                ("p", "Tôi xin cam đoan việc từ chối này là hoàn toàn tự nguyện."),
                ("p", "Người từ chối hưởng di sản"),
                ("p", "LỜI CHỨNG CỦA CÔNG CHỨNG VIÊN"),
                ("p", "Số công chứng 2233.2025/TCDS/CCGD."),
            ],
        )

        payload = extract(docx_path)
        duong_su = payload["web_form"]["duong_su"]
        tai_san = payload["web_form"]["tai_san"]

        self.assertEqual(payload["raw"]["document_kind"], "inheritance_refusal")
        self.assertEqual(payload["web_form"]["ten_hop_dong"], "Văn bản từ chối nhận di sản")
        self.assertEqual(payload["web_form"]["nhom_hop_dong"], "Từ chối nhận di sản thừa kế")
        self.assertEqual(payload["web_form"]["so_cong_chung"], "2233/2025")
        self.assertIn("Hoàng Thị Hợi", payload["web_form"]["nguoi_yeu_cau"])
        self.assertTrue(duong_su.startswith("NGƯỜI TỪ CHỐI NHẬN DI SẢN:"))
        self.assertNotIn("BÊN B:", duong_su)
        self.assertIn("Tài sản thứ hai", tai_san)
        self.assertIn("122(12)", tai_san)
        self.assertNotIn("Bằng Văn bản này, tôi", tai_san)
        self.assertNotIn("LỜI CHỨNG", tai_san)


if __name__ == "__main__":
    unittest.main()
