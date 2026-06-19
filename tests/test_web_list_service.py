from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from ui.services.web_list_service import (
    find_missing_contract_numbers,
    lookup_exported_contract_no,
    read_exported_contract_rows,
)


class WebListServiceTests(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.root = Path(self.tempdir.name)

    def _make_export(self, *values: object) -> Path:
        export_path = self.root / "so_cong_chung.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "SO CONG CHUNG"
        sheet["B1"] = "NGAY"
        for row_index, value in enumerate(values, start=2):
            contract_no = value
            date_value = "01/01/2026"
            if isinstance(value, tuple):
                contract_no, date_value = value
            sheet.cell(row=row_index, column=1).value = contract_no
            sheet.cell(row=row_index, column=2).value = date_value
        workbook.save(export_path)
        workbook.close()
        return export_path

    @staticmethod
    def _date_range() -> dict[str, str]:
        return {"from_date": "01/01/2026", "to_date": "18/06/2026"}

    def test_find_missing_contract_numbers_uses_min_max_by_year(self):
        rows = read_exported_contract_rows(
            self._make_export("405/2026", "407/2026/CCGD", "410/2026"),
            **self._date_range(),
        )

        missing = find_missing_contract_numbers(rows)

        self.assertEqual([item.contract_no for item in missing], ["406/2026", "408/2026", "409/2026"])

    def test_lookup_exported_contract_no_accepts_full_suffix_and_spacing(self):
        rows = read_exported_contract_rows(self._make_export("405/2026"), **self._date_range())

        result = lookup_exported_contract_no(rows, " 405/2026/CCGD ")

        self.assertTrue(result.found)
        self.assertEqual(result.contract_no, "405/2026")
        self.assertEqual(result.row_index, 2)

    def test_read_exported_contract_rows_rejects_short_number_without_slash_year(self):
        export_path = self.root / "short.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "SO CONG CHUNG"
        sheet["B1"] = "NGAY"
        sheet["A2"] = "123"
        sheet["B2"] = "05/02/2026"
        workbook.save(export_path)
        workbook.close()

        rows = read_exported_contract_rows(export_path, **self._date_range())

        self.assertEqual(rows, [])

    def test_web_list_service_does_not_require_folder_or_manifest(self):
        rows = read_exported_contract_rows(
            self._make_export(("100/2025", "01/01/2025"), ("101/2025", "02/01/2025")),
            from_date="01/01/2025",
            to_date="18/06/2025",
        )
        missing = find_missing_contract_numbers(rows)
        result = lookup_exported_contract_no(rows, "100/2025")

        self.assertEqual(missing, [])
        self.assertTrue(result.found)


if __name__ == "__main__":
    unittest.main()
