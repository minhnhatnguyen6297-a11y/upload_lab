from pathlib import Path

from docx import Document
import fitz
from openpyxl import Workbook
import pytest


def _pdf(path: Path, *, with_text: bool) -> Path:
    document = fitz.open()
    page = document.new_page()
    if with_text:
        page.insert_text((72, 72), "Synthetic PDF")
    path.write_bytes(document.tobytes())
    document.close()
    return path


def test_docx_routes_locally_with_plugins_disabled(tmp_path: Path) -> None:
    from poc.conversion_benchmark.router import convert_path

    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Synthetic contract")
    document.save(path)

    envelope = convert_path(path, converter=lambda _: "Synthetic contract")

    assert envelope["route"] == "local"
    assert envelope["content"]["value"] == "Synthetic contract"
    assert envelope["converter"]["config"] == "plugins-disabled"
    assert envelope["converter"]["version"] == "injected-test"


def test_scanned_pdf_is_denied_without_calling_converter(tmp_path: Path) -> None:
    from poc.conversion_benchmark.router import convert_path

    calls = []
    envelope = convert_path(_pdf(tmp_path / "scan.pdf", with_text=False), converter=lambda path: calls.append(path))

    assert envelope["route"] == "ocr_denied"
    assert calls == []
    assert envelope["ocr_calls"] == []
    assert envelope["warnings"] == ["cloud_ocr_not_authorized"]


def test_legacy_doc_and_invalid_file_remain_structured_boundaries(tmp_path: Path) -> None:
    from poc.conversion_benchmark.router import convert_path

    legacy = tmp_path / "legacy.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0synthetic")
    invalid = tmp_path / "broken.pdf"
    invalid.write_bytes(b"%PDF-not-valid")

    legacy_envelope = convert_path(legacy)
    invalid_envelope = convert_path(invalid)

    assert legacy_envelope["route"] == "legacy_doc_external"
    assert "legacy_doc_requires_upload_lab_ifilter" in legacy_envelope["warnings"]
    assert invalid_envelope["route"] == "unsupported"
    assert invalid_envelope["errors"][0]["code"] == "unsupported_input"


def test_real_markitdown_converts_docx_and_xlsx_locally(tmp_path: Path) -> None:
    pytest.importorskip("markitdown")
    from poc.conversion_benchmark.router import convert_path

    docx_path = tmp_path / "contract.docx"
    document = Document()
    document.add_paragraph("Synthetic DOCX content")
    document.save(docx_path)
    xlsx_path = tmp_path / "book.xlsx"
    workbook = Workbook()
    workbook.active.title = "Synthetic"
    workbook.active.append(["Field", "Value"])
    workbook.active.append(["Party", "Synthetic A"])
    workbook.save(xlsx_path)

    docx = convert_path(docx_path)
    xlsx = convert_path(xlsx_path)

    assert "Synthetic DOCX content" in docx["content"]["value"]
    assert "Synthetic" in xlsx["content"]["value"]
    assert docx["ocr_calls"] == xlsx["ocr_calls"] == []


def test_golden_harness_reports_routes_hashes_and_zero_cloud(tmp_path: Path) -> None:
    from poc.conversion_benchmark.harness import materialize_golden, run_benchmark

    report = run_benchmark(
        materialize_golden(tmp_path / "fixtures"),
        tmp_path / "report.json",
        converter=lambda path: "Synthetic PDF text\nSynthetic DOCX paragraph\nSynthetic",
    )

    assert report["summary"] == {"total": 7, "completed": 7, "partial": 0}
    assert report["cloud_call_count"] == 0
    assert report["estimated_cloud_cost_usd"] is None
    assert report["decision"] == "review_required"
    assert all(item["sha256"] and item["route_match"] for item in report["results"])
    assert all(item["provenance_match"] for item in report["results"])
    assert next(item for item in report["results"] if item["sample_id"] == "GD-03")["baseline_text_contains"] is True
    assert next(item for item in report["results"] if item["sample_id"] == "GD-04")["baseline_available"] is True
    assert next(item for item in report["results"] if item["sample_id"] == "GD-01")["recommendation"].startswith("iterate")
