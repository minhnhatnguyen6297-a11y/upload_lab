from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

from extract_contract import extract, read_docx


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT_DIR = BASE_DIR / "regex_review_samples" / "input"
DEFAULT_REPORTS_DIR = BASE_DIR / "regex_review_samples" / "reports"
SUPPORTED_SUFFIXES = {".doc", ".docx"}


def _raw_excerpt(path: Path, *, limit: int = 2500) -> str:
    try:
        text = read_docx(path, use_ifilter_for_doc=path.suffix.lower() == ".doc")
    except Exception as exc:
        return f"[READ_ERROR] {exc}"
    return str(text or "").strip()[:limit]


def _review_file(path: Path, *, root: Path) -> dict:
    relative_name = path.relative_to(root).as_posix()
    item = {
        "file": relative_name,
        "path": str(path),
        "document_kind": "",
        "so_cong_chung": "",
        "nguoi_yeu_cau": "",
        "duong_su": "",
        "tai_san": "",
        "missing_fields": [],
        "raw_excerpt": _raw_excerpt(path),
        "error": "",
    }
    try:
        payload = extract(path)
        web_form = dict(payload.get("web_form") or {})
        raw = dict(payload.get("raw") or {})
        item.update(
            {
                "document_kind": str(raw.get("document_kind") or ""),
                "so_cong_chung": str(web_form.get("so_cong_chung") or ""),
                "nguoi_yeu_cau": str(web_form.get("nguoi_yeu_cau") or ""),
                "duong_su": str(web_form.get("duong_su") or ""),
                "tai_san": str(web_form.get("tai_san") or ""),
                "missing_fields": list(raw.get("missing_web_form_fields") or []),
            }
        )
    except Exception as exc:
        item["error"] = str(exc)
    return item


def _write_csv(path: Path, items: list[dict]) -> None:
    fieldnames = [
        "file",
        "document_kind",
        "so_cong_chung",
        "missing_fields",
        "nguoi_yeu_cau",
        "tai_san",
        "error",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for item in items:
            writer.writerow(
                {
                    "file": item.get("file", ""),
                    "document_kind": item.get("document_kind", ""),
                    "so_cong_chung": item.get("so_cong_chung", ""),
                    "missing_fields": ", ".join(item.get("missing_fields") or []),
                    "nguoi_yeu_cau": item.get("nguoi_yeu_cau", ""),
                    "tai_san": item.get("tai_san", ""),
                    "error": item.get("error", ""),
                }
            )


def _review_flags(item: dict) -> str:
    flags: list[str] = []
    if item.get("error"):
        flags.append("ERROR")
    missing_fields = list(item.get("missing_fields") or [])
    if missing_fields:
        flags.append("MISSING:" + ",".join(missing_fields))
    for field in ("so_cong_chung", "nguoi_yeu_cau", "duong_su", "tai_san"):
        if not str(item.get(field) or "").strip():
            flags.append("BLANK:" + field)
    return " | ".join(flags) if flags else "OK"


def _write_xlsx(path: Path, items: list[dict]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception as exc:  # pragma: no cover - dependency wiring
        raise RuntimeError("Thieu openpyxl trong moi truong tool.") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Regex Review"
    headers = [
        "review_flags",
        "file",
        "document_kind",
        "so_cong_chung",
        "missing_fields",
        "error",
        "nguoi_yeu_cau",
        "duong_su",
        "tai_san",
        "raw_excerpt",
        "path",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    for item in items:
        sheet.append(
            [
                _review_flags(item),
                item.get("file", ""),
                item.get("document_kind", ""),
                item.get("so_cong_chung", ""),
                ", ".join(item.get("missing_fields") or []),
                item.get("error", ""),
                item.get("nguoi_yeu_cau", ""),
                item.get("duong_su", ""),
                item.get("tai_san", ""),
                item.get("raw_excerpt", ""),
                item.get("path", ""),
            ]
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {
        "A": 28,
        "B": 70,
        "C": 22,
        "D": 16,
        "E": 32,
        "F": 40,
        "G": 70,
        "H": 90,
        "I": 90,
        "J": 90,
        "K": 80,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        flag = str(row[0].value or "")
        if flag == "OK":
            row[0].fill = PatternFill("solid", fgColor="E2F0D9")
        else:
            row[0].fill = PatternFill("solid", fgColor="FCE4D6")
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    summary = workbook.create_sheet("Summary")
    total = len(items)
    ok_count = sum(1 for item in items if _review_flags(item) == "OK")
    summary_rows = [
        ("total_files", total),
        ("ok", ok_count),
        ("needs_review", total - ok_count),
        ("errors", sum(1 for item in items if item.get("error"))),
        ("blank_so_cong_chung", sum(1 for item in items if not str(item.get("so_cong_chung") or "").strip())),
        ("blank_nguoi_yeu_cau", sum(1 for item in items if not str(item.get("nguoi_yeu_cau") or "").strip())),
        ("blank_duong_su", sum(1 for item in items if not str(item.get("duong_su") or "").strip())),
        ("blank_tai_san", sum(1 for item in items if not str(item.get("tai_san") or "").strip())),
    ]
    summary.append(("metric", "value"))
    for row in summary_rows:
        summary.append(row)
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 16
    workbook.save(path)


def run_review(
    *,
    input_dir: Path | str = DEFAULT_INPUT_DIR,
    reports_dir: Path | str = DEFAULT_REPORTS_DIR,
) -> Path:
    input_path = Path(input_dir)
    reports_path = Path(reports_dir)
    reports_path.mkdir(parents=True, exist_ok=True)

    files = sorted(
        path
        for path in input_path.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES and not path.name.startswith("~$")
    )
    items = [_review_file(path, root=input_path) for path in files]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "input_dir": str(input_path),
        "files": items,
    }

    json_path = reports_path / f"regex_review_{timestamp}.json"
    csv_path = reports_path / f"regex_review_{timestamp}.csv"
    xlsx_path = reports_path / f"regex_review_{timestamp}.xlsx"
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    _write_csv(csv_path, items)
    _write_xlsx(xlsx_path, items)
    return json_path


def main() -> int:
    report_path = run_review()
    print(f"[OK] Regex review report: {report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
