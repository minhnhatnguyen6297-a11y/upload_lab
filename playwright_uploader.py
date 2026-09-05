from __future__ import annotations

import csv
import importlib
from importlib import util as importlib_util
import io
import json
import os
import re
import shutil
import time
import urllib.error
import urllib.parse
import urllib.request
import warnings
from datetime import date, datetime
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional
import unicodedata

try:
    from dotenv import dotenv_values
except ImportError:  # pragma: no cover
    dotenv_values = None

try:
    from batch_scan import (
        BASE_DIR,
        REGISTRY_DB_PATH,
        connect_registry,
        fetch_registry_records_for_run,
        load_manifest,
        mark_records_uploaded_success,
        now_iso,
        update_registry_record_by_id,
    )
except ImportError:  # pragma: no cover
    from .batch_scan import (
        BASE_DIR,
        REGISTRY_DB_PATH,
        connect_registry,
        fetch_registry_records_for_run,
        load_manifest,
        mark_records_uploaded_success,
        now_iso,
        update_registry_record_by_id,
    )

try:
    from ui.services.environment_check_service import redact_text, safe_url
except ImportError:  # pragma: no cover
    from .ui.services.environment_check_service import redact_text, safe_url

try:
    from uploader_selectors import (
        DEFAULT_GHI_CHU,
        DEFAULT_PHI_CONG_CHUNG,
        DEFAULT_STATUS_TEXT,
        DEFAULT_THU_LAO,
        FORM_FIELD_ORDER,
        FORM_SELECTORS,
        LOGIN_SELECTORS,
        SAVE_BUTTON_SELECTORS,
        VERIFY_FIELDS,
    )
except ImportError:  # pragma: no cover
    from .uploader_selectors import (
        DEFAULT_GHI_CHU,
        DEFAULT_PHI_CONG_CHUNG,
        DEFAULT_STATUS_TEXT,
        DEFAULT_THU_LAO,
        FORM_FIELD_ORDER,
        FORM_SELECTORS,
        LOGIN_SELECTORS,
        SAVE_BUTTON_SELECTORS,
        VERIFY_FIELDS,
    )


QUEUE_STATUSES = ("extracted", "upload_failed", "prepared_dry_run", "prepared_partial")
PREPARE_QUEUE_STATUSES = ("extracted", "upload_failed")
PREPARED_QUEUE_STATUSES = ("prepared_dry_run", "prepared_partial")
ACCESS_TOKEN_STORAGE_KEY = "access_token"
LOGIN_PATH = "/dang-nhap"
CREATE_PATH = "/ho-so-cong-chung/tao-moi-nhanh"
SAVE_API_PATH = "/api/hoso"
MANUAL_LOGIN_TIMEOUT_SECONDS = 300
PLAYWRIGHT_MISSING_MESSAGE = (
    "Chua cai Playwright. Scan/Extract van dung duoc. Muon upload, hay chay "
    "'run_ui.bat' de bootstrap tu dong, hoac cai bang "
    "'pip install -r requirements.txt' va 'playwright install chromium'."
)
UPLOAD_LOG_PATH = BASE_DIR / "logs" / "playwright_uploader.log"
DOWNLOADS_DIR = BASE_DIR / "downloads"
DEFAULT_REQUESTER_SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/1EY5coR_BdMRK-xK6e8d2f1Vj5zr4HbLOqNrd1x89yLk/"
    "edit?gid=1274078459#gid=1274078459"
)
REQUESTER_SHEET_CACHE_NAME = "requester_lookup_cache.csv"
REQUESTER_SHEET_TIMEOUT_SECONDS = 20
STAFF_OPTIONS_CACHE_NAME = "uploader_staff_options.json"
UPLOADER_ENV_KEYS = (
    "ND_BASE_URL",
    "ND_LOGIN_URL",
    "ND_CREATE_URL",
    "ND_STORAGE_STATE_PATH",
    "ND_BROWSER_CHANNEL",
    "ND_MAX_PREPARED_TABS",
    "ND_POST_PREPARE_DELAY_MS",
)
FIELD_VALUE_ALIASES = {
    "nhom_hop_dong": {
        "cam ket - thoa thuan": ["Thoả thuận - Cam kết", "Thỏa thuận - Cam kết", "Cam kết - Thỏa thuận"],
        "thoa thuan - cam ket": ["Thoả thuận - Cam kết", "Thỏa thuận - Cam kết", "Cam kết - Thỏa thuận"],
    }
}


@dataclass
class UploaderSettings:
    base_url: str
    login_url: str
    create_url: str
    storage_state_path: Path
    browser_channel: str = "chromium"
    max_prepared_tabs: int = 10
    post_prepare_delay_ms: int = 1500


@dataclass
class UploadRecord:
    record_id: int
    run_id: str
    contract_no: str
    status: str
    output_json_path: Path
    source_file: Path
    payload: dict
    upload_form: dict
    missing_fields: list[str]
    raw_row: dict


@dataclass
class PreparedBrowserTab:
    record_id: int
    contract_no: str
    page: object
    initial_url: str
    network_log_path: Path
    save_response_ok: bool = False


def _fold_value(text: str) -> str:
    normalized = unicodedata.normalize("NFD", str(text or ""))
    normalized = normalized.replace("\u0111", "d").replace("\u0110", "D")
    folded = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn").lower()
    return " ".join(folded.split())


def get_field_value_candidates(field_name: str, value: str) -> list[str]:
    raw_value = str(value or "").strip()
    if not raw_value:
        return []

    candidates: list[str] = [raw_value]
    alias_map = FIELD_VALUE_ALIASES.get(field_name, {})
    for alias in alias_map.get(_fold_value(raw_value), []):
        if alias and alias not in candidates:
            candidates.append(alias)
    return candidates


def field_value_matches(field_name: str, expected: str, actual: str) -> bool:
    actual_fold = _fold_value(actual)
    if not actual_fold:
        return False

    for candidate in get_field_value_candidates(field_name, expected):
        candidate_fold = _fold_value(candidate)
        if field_name in {"cong_chung_vien", "thu_ky"}:
            if candidate_fold and candidate_fold == actual_fold:
                return True
        elif candidate_fold and candidate_fold in actual_fold:
            return True
    return False


def _default_log(message: str) -> None:
    try:
        print(message)
    except UnicodeEncodeError:
        # Some Windows terminals still expose a CP1252 stdout stream. Keep
        # the diagnostic in the UTF-8 file log while ensuring a Playwright
        # launch/navigation error can still be returned to the caller.
        print(str(message).encode("ascii", errors="replace").decode("ascii"))


def _default_uploader_env_values(base_dir: Path = BASE_DIR) -> dict[str, str]:
    base_url = "https://congchungnamdinh.ninhbinh.gov.vn"
    return {
        "ND_BASE_URL": base_url,
        "ND_LOGIN_URL": "",
        "ND_CREATE_URL": "",
        "ND_STORAGE_STATE_PATH": "nd_storage_state.json",
        "ND_BROWSER_CHANNEL": "chromium",
        "ND_MAX_PREPARED_TABS": "10",
        "ND_POST_PREPARE_DELAY_MS": "1500",
    }


def _normalize_env_value(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if len(text) >= 2 and text[0] == text[-1] == '"':
        return text[1:-1].replace('\\"', '"')
    return text


def _resolve_tool_relative_path(base_dir: Path, raw_path: str, *, default_name: str) -> Path:
    raw = str(raw_path or "").strip() or default_name
    normalized = raw.replace("\\", "/")
    if normalized.startswith("tools/upload_lab/"):
        normalized = normalized[len("tools/upload_lab/") :]
    if normalized.startswith("UPLOAD/"):
        normalized = normalized[len("UPLOAD/") :]
    if normalized.startswith("upload/"):
        normalized = normalized[len("upload/") :]
    path = Path(normalized)
    if path.is_absolute():
        return path
    return (base_dir / path).resolve()


def ensure_uploader_env_file(base_dir: Path = BASE_DIR) -> Path:
    base_dir = Path(base_dir)
    env_path = base_dir / ".env"
    env_example_path = base_dir / ".env.example"
    if env_path.exists():
        return env_path

    if env_example_path.exists():
        shutil.copyfile(env_example_path, env_path)
    else:
        save_uploader_env(_default_uploader_env_values(base_dir), base_dir=base_dir)
    return env_path


def read_uploader_env(base_dir: Path = BASE_DIR, *, ensure_exists: bool = False) -> dict[str, str]:
    base_dir = Path(base_dir)
    if ensure_exists:
        ensure_uploader_env_file(base_dir)

    values = _default_uploader_env_values(base_dir)
    env_path = base_dir / ".env"
    if env_path.exists():
        if dotenv_values is not None:
            loaded = dotenv_values(env_path)
            for key, value in loaded.items():
                if key:
                    values[key] = _normalize_env_value(value)
        else:
            for line in env_path.read_text(encoding="utf-8").splitlines():
                stripped = line.strip()
                if not stripped or stripped.startswith("#") or "=" not in stripped:
                    continue
                key, value = stripped.split("=", 1)
                values[key.strip()] = _normalize_env_value(value)

    for key in UPLOADER_ENV_KEYS:
        if key in os.environ:
            values[key] = _normalize_env_value(os.environ[key])
    return values


def _format_env_value(value: object) -> str:
    text = str(value or "")
    if any(ch.isspace() for ch in text) or "#" in text or '"' in text:
        return '"' + text.replace("\\", "\\\\").replace('"', '\\"') + '"'
    return text


def save_uploader_env(values: dict[str, object], *, base_dir: Path = BASE_DIR) -> Path:
    base_dir = Path(base_dir)
    env_path = base_dir / ".env"
    merged = _default_uploader_env_values(base_dir)
    merged.update({key: _normalize_env_value(value) for key, value in values.items() if key in UPLOADER_ENV_KEYS})

    lines = ["# Cau hinh uploader cho tool standalone", ""]
    for key in UPLOADER_ENV_KEYS:
        lines.append(f"{key}={_format_env_value(merged.get(key, ''))}")
    env_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return env_path


def update_uploader_env(values: dict[str, object], *, base_dir: Path = BASE_DIR) -> Path:
    """Update selected settings while stripping legacy username/password keys."""
    merged = {
        key: value
        for key, value in read_uploader_env(base_dir, ensure_exists=True).items()
        if key in UPLOADER_ENV_KEYS
    }
    merged.update({key: value for key, value in values.items() if key in UPLOADER_ENV_KEYS})
    return save_uploader_env(merged, base_dir=base_dir)


def get_uploader_setup_status(base_dir: Path = BASE_DIR) -> dict[str, object]:
    base_dir = Path(base_dir)
    env_path = base_dir / ".env"
    values = read_uploader_env(base_dir)
    storage_state_path = _resolve_tool_relative_path(
        base_dir,
        values.get("ND_STORAGE_STATE_PATH", ""),
        default_name="nd_storage_state.json",
    )
    missing_fields = [key for key in ("ND_BASE_URL",) if not str(values.get(key) or "").strip()]
    storage_state_exists = storage_state_path.exists()
    ready = not missing_fields and storage_state_exists
    if ready:
        message = f"Uploader da san sang. Storage state: {storage_state_path.name}"
    elif missing_fields:
        labels = {
            "ND_BASE_URL": "base url",
        }
        missing_text = ", ".join(labels[key] for key in missing_fields)
        message = f"Can cau hinh uploader lan dau ({missing_text}). Bam 'Cau hinh'."
    else:
        message = "Can dang nhap uploader lan dau de tao nd_storage_state.json. Bam 'Cau hinh'."
    return {
        "env_path": env_path,
        "env_exists": env_path.exists(),
        "values": values,
        "missing_fields": missing_fields,
        "storage_state_path": storage_state_path,
        "storage_state_exists": storage_state_exists,
        "ready": ready,
        "message": message,
    }


def append_log_line(path: Path, message: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    line = f"{now_iso()} {message}\n"
    with path.open("a", encoding="utf-8") as handle:
        handle.write(line)


def probe_playwright_runtime() -> tuple[bool, str]:
    if importlib_util.find_spec("playwright") is None:
        return False, PLAYWRIGHT_MISSING_MESSAGE

    try:
        importlib.import_module("playwright.sync_api")
    except Exception as exc:
        return False, f"Playwright chua san sang: {exc}"

    return True, "Playwright package san sang cho dry-run upload."


def sanitize_contract_no(value: str) -> str:
    text = str(value or "").strip()
    return text.replace("/", "_").replace("\\", "_")


def normalize_web_contract_no(value: str) -> str:
    text = re.sub(r"\s+", "", str(value or "").strip().upper())
    match = re.match(r"^(\d+)[./](\d{4})(?:/[A-Z0-9]+)*(?:/CCGD)?$", text)
    if match:
        return f"{match.group(1)}/{match.group(2)}"
    if text.endswith("/CCGD"):
        text = text[:-5]
    return text


def normalize_contract_no_for_compare(value: str) -> str:
    text = normalize_web_contract_no(value)
    return re.sub(r"\s+", "", text)


def normalize_requester_sheet_csv_url(sheet_url: str) -> str:
    raw_url = str(sheet_url or "").strip()
    if not raw_url:
        return ""

    parsed = urllib.parse.urlparse(raw_url)
    if not parsed.scheme or not parsed.netloc:
        raise ValueError("Link doi chieu phai la link Google Sheet hop le.")

    host = parsed.netloc.lower()
    if "docs.google.com" not in host:
        return raw_url

    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", parsed.path)
    if not match:
        raise ValueError("Khong doc duoc ma Google Sheet tu link doi chieu.")

    query = urllib.parse.parse_qs(parsed.query or "")
    gid = str(query.get("gid", [""])[0] or "").strip()
    if not gid and str(parsed.fragment or "").startswith("gid="):
        gid = parsed.fragment.split("=", 1)[1].strip()
    gid = gid or "0"
    return f"https://docs.google.com/spreadsheets/d/{match.group(1)}/gviz/tq?tqx=out:csv&gid={gid}"


def _extract_lookup_year(value: str) -> str:
    text = str(value or "").strip()
    match = re.search(r"(20\d{2})", text)
    return match.group(1) if match else ""


def build_requester_lookup_keys(contract_no: str, *, date_text: str = "") -> list[str]:
    text = str(contract_no or "").strip()
    match = re.search(r"(\d+)", text)
    if not match:
        return []

    number_key = str(int(match.group(1)))
    year = ""
    year_match = re.search(r"/(20\d{2})(?:\D|$)", text)
    if year_match:
        year = year_match.group(1)
    if not year:
        year = _extract_lookup_year(date_text)

    keys: list[str] = []
    if year:
        keys.append(f"{number_key}/{year}")
    keys.append(number_key)
    return keys


def parse_requester_sheet_csv(csv_text: str) -> dict[str, str]:
    content = str(csv_text or "").strip()
    if not content:
        return {}

    reader = csv.DictReader(io.StringIO(content))
    fieldnames = list(reader.fieldnames or [])
    if not fieldnames:
        return {}

    def pick_field(*candidates: str) -> str:
        for field in fieldnames:
            folded = _fold_value(field)
            if any(candidate in folded for candidate in candidates):
                return field
        return ""

    contract_field = pick_field("so cc", "so cong chung")
    requester_field = pick_field("ten khach hang", "nguoi yeu cau")
    date_field = pick_field("ngay cc", "ngay cong chung")
    if not contract_field or not requester_field:
        return {}

    lookup: dict[str, str] = {}
    for row in reader:
        contract_value = str(row.get(contract_field) or "").strip()
        requester_name = str(row.get(requester_field) or "").strip()
        date_value = str(row.get(date_field) or "").strip()
        if not contract_value or not requester_name:
            continue
        for key in build_requester_lookup_keys(contract_value, date_text=date_value):
            lookup.setdefault(key, requester_name)
    return lookup


def load_requester_contract_lookup(
    *,
    working_dir: Path = BASE_DIR,
    sheet_url: str = "",
    log_callback: Optional[Callable[[str], None]] = None,
) -> tuple[dict[str, str], str | None]:
    requested_url = str(sheet_url or "").strip()
    if not requested_url:
        return {}, None

    try:
        csv_url = normalize_requester_sheet_csv_url(requested_url)
    except ValueError as exc:
        return {}, str(exc)

    cache_path = Path(working_dir) / DOWNLOADS_DIR.name / REQUESTER_SHEET_CACHE_NAME
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    logger = log_callback or _default_log
    csv_text = ""
    warning_text: str | None = None

    request = urllib.request.Request(csv_url, headers={"User-Agent": "upload-lab/1.0"})
    try:
        with urllib.request.urlopen(request, timeout=REQUESTER_SHEET_TIMEOUT_SECONDS) as response:
            csv_bytes = response.read()
        csv_text = csv_bytes.decode("utf-8-sig", errors="replace")
        cache_path.write_text(csv_text, encoding="utf-8")
        logger(f"[UPLOAD][REQUESTER] Tai lookup nguoi yeu cau OK: {csv_url}")
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        if cache_path.exists():
            csv_text = cache_path.read_text(encoding="utf-8")
            warning_text = "Khong tai duoc link doi chieu. Dang dung du lieu cache gan nhat."
            logger(f"[UPLOAD][REQUESTER] Loi tai link doi chieu, dung cache {cache_path.name}: {exc}")
        else:
            warning_text = "Khong tai duoc link doi chieu nguoi yeu cau cong chung."
            logger(f"[UPLOAD][REQUESTER] Loi tai link doi chieu: {exc}")
            return {}, warning_text

    lookup = parse_requester_sheet_csv(csv_text)
    if not lookup:
        return {}, warning_text or "Link doi chieu khong co du lieu hop le."

    logger(f"[UPLOAD][REQUESTER] Da nap {len(lookup)} khoa doi chieu nguoi yeu cau.")
    return lookup, warning_text


def _normalize_person_name(value: str) -> str:
    folded = _fold_value(value)
    return re.sub(r"^(ong|ba|anh|chi)\s+", "", folded).strip()


def _requester_name_matches(expected_name: str, actual_name: str) -> bool:
    expected = _normalize_person_name(expected_name)
    actual = _normalize_person_name(actual_name)
    if not expected or not actual:
        return False
    return expected == actual or expected in actual or actual in expected


def _format_requester_person(person: dict) -> str:
    title = str(person.get("gioi_tinh") or "").strip()
    name = str(person.get("ho_ten") or person.get("name") or "").strip()
    first_line = " ".join(part for part in (title, name) if part).strip() or name
    lines: list[str] = []
    if first_line:
        birthday = str(person.get("ngay_sinh") or "").strip()
        if birthday:
            first_line += f", sinh ngay: {birthday}"
        lines.append(first_line)

    cccd = str(person.get("cccd") or "").strip()
    if cccd:
        id_line = f"Can cuoc so: {cccd}"
        issue_place = str(person.get("noi_cap") or "").strip()
        issue_date = str(person.get("ngay_cap_cccd") or "").strip()
        if issue_place and issue_date:
            id_line += f" do {issue_place} cap ngay {issue_date}"
        elif issue_place:
            id_line += f" do {issue_place} cap"
        elif issue_date:
            id_line += f" cap ngay {issue_date}"
        lines.append(id_line)

    address_line = str(person.get("dia_chi_line") or "").strip(" ;")
    if not address_line:
        address_value = str(person.get("dia_chi") or "").strip(" ;")
        if address_value:
            address_line = f"Thuong tru tai: {address_value}"
    if address_line:
        lines.append(address_line)
    return "\n".join(line for line in lines if line).strip()


def resolve_requester_override(
    payload: dict,
    *,
    contract_no: str,
    requester_lookup: Optional[dict[str, str]] = None,
) -> str:
    lookup = dict(requester_lookup or {})
    if not lookup:
        return ""

    web_form = dict(payload.get("web_form") or {})
    raw = dict(payload.get("raw") or {})
    date_text = str(web_form.get("ngay_cong_chung") or raw.get("ngay_cong_chung") or "").strip()
    lookup_keys = build_requester_lookup_keys(contract_no or web_form.get("so_cong_chung", ""), date_text=date_text)
    requester_name = next((lookup.get(key, "") for key in lookup_keys if lookup.get(key, "")), "")
    requester_name = str(requester_name or "").strip()
    if not requester_name:
        return ""

    for party_name in ("ben_a", "ben_b"):
        party = dict(raw.get(party_name) or {})
        for person in list(party.get("nguoi") or []):
            person_name = str(person.get("ho_ten") or person.get("name") or person.get("raw_text") or "").strip()
            if _requester_name_matches(requester_name, person_name):
                formatted = _format_requester_person(person)
                if formatted:
                    return formatted
    return requester_name


def read_exported_contract_numbers(export_path: Path | str) -> set[str]:
    path = Path(export_path)
    if not path.exists():
        raise FileNotFoundError(f"Khong tim thay file excel doi chieu: {path}")

    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - dependency wiring
        raise RuntimeError(
            "Thieu openpyxl. Hay chay lai bootstrap hoac cai 'openpyxl' trong moi truong cua tool."
        ) from exc

    contract_nos: set[str] = set()
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style, apply openpyxl's default",
            category=UserWarning,
        )
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        for row_index, (cell_value,) in enumerate(
            sheet.iter_rows(min_col=1, max_col=1, values_only=True),
            start=1,
        ):
            raw = str(cell_value or "").strip()
            if not raw:
                continue
            normalized = normalize_contract_no_for_compare(raw)
            header_key = re.sub(r"[^A-Z]", "", _fold_value(raw).upper())
            if row_index == 1 and "SOCONGCHUNG" in header_key:
                continue
            if not normalized:
                continue
            contract_nos.add(normalized)
    finally:
        workbook.close()

    return contract_nos


def split_records_by_existing_contract_nos(
    records: list["UploadRecord"],
    existing_contract_nos: set[str],
) -> tuple[list["UploadRecord"], list["UploadRecord"]]:
    if not existing_contract_nos:
        return list(records), []

    filtered: list[UploadRecord] = []
    duplicates: list[UploadRecord] = []
    for record in records:
        normalized = normalize_contract_no_for_compare(record.contract_no or record.upload_form.get("so_cong_chung", ""))
        if normalized and normalized in existing_contract_nos:
            duplicates.append(record)
        else:
            filtered.append(record)
    return filtered, duplicates


def default_export_from_date() -> str:
    return os.getenv("ND_WEB_EXPORT_FROM_DATE", "01/01/2026").strip() or "01/01/2026"


def default_export_to_date() -> str:
    return date.today().strftime("%d/%m/%Y")


def _parse_export_date_input(value: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        raise ValueError("Ngay xuat so cong chung khong duoc de trong.")

    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, fmt).strftime("%d/%m/%Y")
        except ValueError:
            continue
    raise ValueError(f"Ngay khong hop le: {raw}. Dung DD/MM/YYYY hoac YYYY-MM-DD.")


def load_uploader_settings(base_dir: Path = BASE_DIR) -> UploaderSettings:
    values = read_uploader_env(base_dir)

    base_url = str(values.get("ND_BASE_URL") or "https://congchungnamdinh.ninhbinh.gov.vn").rstrip("/")
    login_url = str(values.get("ND_LOGIN_URL") or f"{base_url}{LOGIN_PATH}").strip()
    create_url = str(values.get("ND_CREATE_URL") or f"{base_url}{CREATE_PATH}").strip()
    storage_state_path = _resolve_tool_relative_path(
        Path(base_dir),
        str(values.get("ND_STORAGE_STATE_PATH") or ""),
        default_name="nd_storage_state.json",
    )
    browser_channel = str(values.get("ND_BROWSER_CHANNEL") or "chromium").strip().lower() or "chromium"
    max_tabs = int(str(values.get("ND_MAX_PREPARED_TABS") or "10"))
    delay_ms = int(str(values.get("ND_POST_PREPARE_DELAY_MS") or "1500"))

    return UploaderSettings(
        base_url=base_url,
        login_url=login_url,
        create_url=create_url,
        storage_state_path=storage_state_path,
        browser_channel=browser_channel,
        max_prepared_tabs=max_tabs,
        post_prepare_delay_ms=delay_ms,
    )


def load_upload_queue(
    manifest_path: Path | str,
    *,
    working_dir: Path = BASE_DIR,
    limit: Optional[int] = None,
    selected_record_ids: set[int] | None = None,
    requester_lookup: Optional[dict[str, str]] = None,
    cong_chung_vien: str | None = None,
    thu_ky: str | None = None,
    statuses: tuple[str, ...] = QUEUE_STATUSES,
) -> tuple[dict, list[UploadRecord], int]:
    manifest = load_manifest(manifest_path)
    run_id = str(manifest.get("run_id") or "").strip()
    if not run_id:
        raise ValueError("Manifest khong co run_id")

    conn = connect_registry(working_dir / REGISTRY_DB_PATH.name)
    try:
        rows = fetch_registry_records_for_run(conn, run_id, statuses=statuses)
    finally:
        conn.close()

    records: list[UploadRecord] = []
    for row in rows:
        row_dict = dict(row)
        output_json_path = Path(row_dict.get("output_json_path") or "")
        if not output_json_path.exists():
            continue
        payload = json.loads(output_json_path.read_text(encoding="utf-8"))
        upload_form = build_upload_form_data(
            payload,
            contract_no=str(row_dict.get("contract_no") or ""),
            requester_lookup=requester_lookup,
            cong_chung_vien=cong_chung_vien,
            thu_ky=thu_ky,
        )
        source_file = Path(upload_form["file_hop_dong"])
        missing_fields = identify_missing_fields(upload_form)
        records.append(
            UploadRecord(
                record_id=int(row_dict["id"]),
                run_id=run_id,
                contract_no=str(row_dict.get("contract_no") or upload_form.get("so_cong_chung") or ""),
                status=str(row_dict.get("status") or ""),
                output_json_path=output_json_path,
                source_file=source_file,
                payload=payload,
                upload_form=upload_form,
                missing_fields=missing_fields,
                raw_row=row_dict,
            )
        )

    total_pending = len(records)
    if selected_record_ids is not None:
        selected = {int(record_id) for record_id in selected_record_ids}
        records = [record for record in records if int(record.record_id) in selected]
    if limit is not None:
        records = records[:limit]
    return manifest, records, total_pending


def finalize_uploaded_records(record_ids: list[int], *, working_dir: Path = BASE_DIR) -> int:
    if not record_ids:
        return 0
    conn = connect_registry(working_dir / REGISTRY_DB_PATH.name)
    try:
        mark_records_uploaded_success(conn, [int(rid) for rid in record_ids])
    finally:
        conn.close()
    return len(record_ids)


def build_upload_form_data(
    payload: dict,
    *,
    contract_no: str = "",
    requester_lookup: Optional[dict[str, str]] = None,
    cong_chung_vien: str | None = None,
    thu_ky: str | None = None,
) -> dict:
    web_form = dict(payload.get("web_form") or {})
    raw = dict(payload.get("raw") or {})
    normalized_contract_no = normalize_web_contract_no(web_form.get("so_cong_chung") or contract_no or "")
    return {
        "ten_hop_dong": str(web_form.get("ten_hop_dong") or "").strip(),
        "ngay_cong_chung": str(web_form.get("ngay_cong_chung") or "").strip(),
        "so_cong_chung": normalized_contract_no,
        "tinh_trang": DEFAULT_STATUS_TEXT,
        "nhom_hop_dong": str(web_form.get("nhom_hop_dong") or "").strip(),
        "loai_tai_san": str(web_form.get("loai_tai_san") or "").strip(),
        "cong_chung_vien": str(
            cong_chung_vien if cong_chung_vien is not None else web_form.get("cong_chung_vien") or ""
        ).strip(),
        "thu_ky": str(thu_ky if thu_ky is not None else web_form.get("thu_ky") or "").strip(),
        "nguoi_yeu_cau": (
            resolve_requester_override(
                payload,
                contract_no=normalized_contract_no or str(contract_no or "").strip(),
                requester_lookup=requester_lookup,
            )
            or str(web_form.get("nguoi_yeu_cau") or "").strip()
        ),
        "duong_su": str(web_form.get("duong_su") or "").strip(),
        "tai_san": str(web_form.get("tai_san") or "").strip(),
        "ghi_chu": DEFAULT_GHI_CHU,
        "phi_cong_chung": DEFAULT_PHI_CONG_CHUNG,
        "thu_lao_cong_chung": DEFAULT_THU_LAO,
        "file_hop_dong": str(raw.get("file_goc") or "").strip(),
    }


def identify_missing_fields(upload_form: dict) -> list[str]:
    critical_fields = ("ten_hop_dong", "so_cong_chung", "nhom_hop_dong", "loai_tai_san", "tai_san")
    return [field for field in critical_fields if not str(upload_form.get(field) or "").strip()]


class NamDinhUploaderSession:
    def __init__(
        self,
        settings: UploaderSettings,
        *,
        working_dir: Path = BASE_DIR,
        log_callback: Optional[Callable[[str], None]] = None,
    ):
        self.settings = settings
        self.working_dir = Path(working_dir)
        self._log_callback = log_callback or _default_log
        self.log_path = self.working_dir / "logs" / "playwright_uploader.log"
        self.run_log_path: Path | None = None
        self._playwright = None
        self.browser = None
        self.context = None
        self.anchor_page = None
        self.login_page = None
        self.login_started_at = 0.0
        self.login_response_ok = False
        self.prepared_pages: dict[int, PreparedBrowserTab] = {}
        self._prepared_record_ids: set[int] = set()

    def log(self, message: str) -> None:
        append_log_line(self.log_path, message)
        if self.run_log_path is not None:
            append_log_line(self.run_log_path, message)
        self._log_callback(message)

    @staticmethod
    def _literal_xpath(text: str) -> str:
        if "'" not in text:
            return f"'{text}'"
        if '"' not in text:
            return f'"{text}"'
        parts = text.split("'")
        return "concat(" + ", \"'\", ".join(f"'{part}'" for part in parts) + ")"

    def close(self) -> None:
        self.anchor_page = None
        self.login_page = None
        self.login_started_at = 0.0
        self.login_response_ok = False
        self.prepared_pages.clear()
        self._prepared_record_ids.clear()
        if self.context is not None:
            try:
                self.context.close()
            except Exception:
                pass
            self.context = None
        if self.browser is not None:
            try:
                self.browser.close()
            except Exception:
                pass
            self.browser = None
        if self._playwright is not None:
            try:
                self._playwright.stop()
            except Exception:
                pass
            self._playwright = None

    def _import_playwright(self):
        try:
            from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
            from playwright.sync_api import sync_playwright
        except Exception as exc:
            raise RuntimeError(PLAYWRIGHT_MISSING_MESSAGE) from exc
        return sync_playwright, PlaywrightTimeoutError

    def _ensure_context(self) -> None:
        if self.context is not None and self.browser is not None:
            try:
                _ = self.context.pages
                return
            except Exception:
                self.close()

        sync_playwright, _ = self._import_playwright()
        if self._playwright is None:
            self._playwright = sync_playwright().start()

        launch_kwargs = {"headless": False, "args": ["--start-maximized"]}
        if self.settings.browser_channel and self.settings.browser_channel not in {"", "chromium"}:
            launch_kwargs["channel"] = self.settings.browser_channel

        self.browser = self._playwright.chromium.launch(**launch_kwargs)
        self.settings.storage_state_path.parent.mkdir(parents=True, exist_ok=True)
        context_kwargs = {"accept_downloads": True, "viewport": None, "no_viewport": True}
        if self.settings.storage_state_path.exists():
            context_kwargs["storage_state"] = str(self.settings.storage_state_path)
        try:
            self.context = self.browser.new_context(**context_kwargs)
        except Exception:
            # A partially written or old storage-state file must not prevent login.
            context_kwargs.pop("storage_state", None)
            self.context = self.browser.new_context(**context_kwargs)

    @staticmethod
    def _classify_navigation_error(exc: object) -> str:
        text = str(exc or "").lower()
        if "timeout" in text:
            return "timeout"
        if "certificate" in text or "cert_" in text:
            return "certificate"
        if "dns" in text or "name_not_resolved" in text or "err_name_not_resolved" in text:
            return "dns"
        if "proxy" in text or "access_denied" in text or "blocked" in text:
            return "proxy_or_policy"
        if any(
            marker in text
            for marker in (
                "browsertype.launch",
                "browser launch",
                "executable doesn't exist",
                "executable does not exist",
                "missing executable",
                "playwright install",
            )
        ):
            return "browser_launch"
        return "navigation_error"

    def preflight_login(self, *, timeout_ms: int = 15000) -> dict:
        """Launch/probe the login page while retaining the same browser session.

        This is intentionally a method on the existing uploader session.  The
        caller can therefore continue manual login on ``login_page`` without
        opening a second Chromium instance after the environment check.
        """

        target_url = self.settings.login_url or self.settings.base_url
        clean_url = safe_url(target_url)
        if not clean_url:
            return {
                "status": "blocked",
                "error_code": "invalid_url",
                "message": "Địa chỉ trang đăng nhập chưa hợp lệ.",
                "guidance": "Mở Cấu hình và nhập địa chỉ web HTTP/HTTPS hợp lệ.",
                "url": "",
                "browser_channel": self.settings.browser_channel,
                "reused_browser": False,
            }

        try:
            self._ensure_context()
            page = self.anchor_page
            reused_browser = not self._page_is_closed(page)
            if self._page_is_closed(page):
                page = self.context.new_page()
            self.anchor_page = page
            page.goto(clean_url, wait_until="domcontentloaded", timeout=int(timeout_ms))
            self._begin_login_tracking(page)
            self.log("[ENV] Chromium va trang dang nhap san sang.")
            return {
                "status": "passed",
                "error_code": "",
                "message": "Chromium đã mở và truy cập được trang đăng nhập.",
                "guidance": "Hãy đăng nhập trên Chromium đang mở; app sẽ tự nhận diện.",
                "url": safe_url(str(getattr(page, "url", "") or target_url)),
                "browser_channel": self.settings.browser_channel,
                "reused_browser": reused_browser,
            }
        except Exception as exc:
            error_code = self._classify_navigation_error(exc)
            messages = {
                "timeout": "Trang đăng nhập phản hồi quá thời gian cho phép.",
                "certificate": "Kết nối tới trang đăng nhập gặp lỗi chứng chỉ.",
                "dns": "Không phân giải được tên miền trang đăng nhập.",
                "proxy_or_policy": "Proxy hoặc chính sách mạng chặn trang đăng nhập.",
                "browser_launch": "Không thể khởi động Chromium của app.",
                "navigation_error": "Không thể mở trang đăng nhập bằng Chromium.",
            }
            self.log(f"[ENV] {messages[error_code]} ({redact_text(exc)})")
            return {
                "status": "blocked",
                "error_code": error_code,
                "message": messages[error_code],
                "guidance": "Kiểm tra mạng, VPN, proxy, chứng chỉ hoặc cài đặt Chromium rồi thử lại.",
                "url": safe_url(str(getattr(self.anchor_page, "url", "") or target_url)),
                "browser_channel": self.settings.browser_channel,
                "reused_browser": not self._page_is_closed(self.anchor_page),
            }

    def _save_storage_state(self) -> None:
        if self.context is not None:
            self.context.storage_state(path=str(self.settings.storage_state_path))

    @staticmethod
    def _page_is_closed(page) -> bool:
        if page is None:
            return True
        try:
            return bool(page.is_closed())
        except Exception:
            return False

    @staticmethod
    def _clean_network_url(url: str) -> str:
        parsed = urllib.parse.urlsplit(str(url or ""))
        return urllib.parse.urlunsplit((parsed.scheme, parsed.netloc, parsed.path, "", ""))

    def _append_network_event(
        self,
        path: Path,
        *,
        method: str,
        url: str,
        status: int,
        record_id: int | None = None,
    ) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        payload = {
            "time": now_iso(),
            "method": str(method or "").upper(),
            "url": self._clean_network_url(url),
            "status": int(status or 0),
        }
        if record_id is not None:
            payload["record_id"] = int(record_id)
        with path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(payload, ensure_ascii=False) + "\n")

    def _access_token(self, page) -> str:
        try:
            return str(
                page.evaluate(
                    f"() => window.localStorage.getItem({json.dumps(ACCESS_TOKEN_STORAGE_KEY)}) || ''"
                )
                or ""
            ).strip()
        except Exception:
            return ""

    def _wait_for_create_or_login(self, page, *, timeout_ms: int = 30000) -> str:
        if not hasattr(page, "locator"):
            # Lightweight fakes used by unit tests expose no DOM API.
            return "create"
        deadline = time.monotonic() + max(int(timeout_ms), 0) / 1000
        while time.monotonic() <= deadline:
            if self._is_login_page(page):
                return "login"
            for field_name in ("so_cong_chung", "ten_hop_dong"):
                try:
                    locator = self._find_control_locator(page, field_name)
                    if locator is not None and locator.count() > 0:
                        return "create"
                except Exception:
                    continue
            page.wait_for_timeout(250)
        raise RuntimeError("Trang tao ho so tai qua lau hoac khong dung cau truc mong doi.")

    def _begin_login_tracking(self, page) -> None:
        self.login_page = page
        self.anchor_page = page
        self.login_started_at = time.monotonic()
        self.login_response_ok = False
        network_log_path = self.working_dir / "logs" / "login_network.jsonl"

        def handle_response(response) -> None:
            try:
                request = response.request
                method = str(request.method or "").upper()
                url = str(response.url or "")
                status = int(response.status or 0)
                path = urllib.parse.urlsplit(url).path.rstrip("/")
                if method == "POST" and path == "/api/auth/login":
                    self._append_network_event(
                        network_log_path,
                        method=method,
                        url=url,
                        status=status,
                    )
                    if 200 <= status < 300:
                        self.login_response_ok = True
            except Exception:
                pass

        try:
            page.on("response", handle_response)
        except Exception:
            pass

    def open_manual_login(self) -> dict:
        self._ensure_context()
        if not self._page_is_closed(self.login_page):
            self.log("[LOGIN] Chromium dang mo san trang dang nhap.")
            return {
                "status": "waiting",
                "message": "Chromium đang mở sẵn trang đăng nhập; hãy đăng nhập trên trình duyệt.",
                "url": safe_url(str(getattr(self.login_page, "url", "") or self.settings.login_url)),
            }
        page = self.anchor_page
        if self._page_is_closed(page):
            page = self.context.new_page()
        self._begin_login_tracking(page)
        # Chi mo dung dia chi web mac dinh, khong tu dieu huong vao trang nhap ho so.
        target_url = self.settings.base_url or self.settings.login_url
        try:
            page.goto(target_url, wait_until="domcontentloaded")
        except Exception as exc:
            # Chromium has launched successfully. Keep it open so the user can
            # retry in the address bar instead of losing the whole app flow.
            message = f"Khong mo duoc trang dang nhap: {exc}"
            self.log(f"[LOGIN] {message}")
            return {
                "status": "navigation_error",
                "message": message,
                "url": str(getattr(page, "url", "") or target_url),
            }
        self.log("[LOGIN] Da mo Chromium. Hay dang nhap tren trinh duyet; app se tu nhan dien.")
        return {"status": "waiting", "url": str(page.url or target_url)}

    def poll_manual_login(self, *, force: bool = False) -> dict:
        page = self.login_page
        if page is None:
            return {"status": "idle", "message": "Chua mo trang dang nhap."}
        if self._page_is_closed(page):
            self.login_page = None
            self.login_started_at = 0.0
            return {"status": "closed", "message": "Cua so dang nhap da bi dong."}

        token = self._access_token(page)
        url = str(getattr(page, "url", "") or "")
        logged_in = bool(token) and not self._is_login_page(page)
        if logged_in:
            self._save_storage_state()
            self.login_page = None
            self.login_started_at = 0.0
            self.login_response_ok = False
            self.anchor_page = page
            self.log("[LOGIN] Da nhan dien dang nhap thanh cong va luu session.")
            return {"status": "authenticated", "message": "Da dang nhap.", "url": url}

        elapsed = time.monotonic() - self.login_started_at if self.login_started_at else 0.0
        if elapsed >= MANUAL_LOGIN_TIMEOUT_SECONDS:
            self.login_started_at = 0.0
            return {
                "status": "timeout",
                "message": "Chua nhan dien duoc dang nhap sau 5 phut. Trinh duyet van duoc giu mo.",
                "url": url,
            }
        message = "Dang cho ban dang nhap tren Chromium."
        if force:
            message = "Chua thay token dang nhap. Hay dang nhap xong va doi trang roi thu lai."
        return {"status": "waiting", "message": message, "url": url}

    def wait_for_manual_login(self, page, *, stop_event=None) -> None:
        if self.login_page is not page:
            self._begin_login_tracking(page)
        while True:
            if stop_event is not None and stop_event.is_set():
                raise RuntimeError("Da dung trong luc cho dang nhap.")
            result = self.poll_manual_login()
            if result["status"] == "authenticated":
                return
            if result["status"] in {"closed", "timeout"}:
                raise RuntimeError(str(result.get("message") or "Khong the dang nhap."))
            page.wait_for_timeout(500)

    def _locator_from_strategy(self, page, strategy: dict, *, dynamic_text: str = ""):
        stype = strategy.get("type")
        if stype == "css":
            return page.locator(strategy["value"]).first
        if stype == "xpath":
            return page.locator(strategy["value"]).first
        if stype == "text":
            return page.get_by_text(strategy["value"], exact=False).first
        if stype == "role":
            return page.get_by_role(strategy["role"], name=strategy.get("name"), exact=False).first
        if stype == "text_dynamic":
            if dynamic_text:
                return page.get_by_text(dynamic_text, exact=False).first
        return None

    def _label_locator_xpath(self, label_text: str, suffix_xpath: str):
        literal = self._literal_xpath(label_text)
        return (
            f"xpath=((//label[contains(normalize-space(.), {literal})])[1])/{suffix_xpath}"
        )

    def _find_label_locator(self, page, label_text: str):
        literal = self._literal_xpath(label_text)
        candidates = [
            f"xpath=(//label[contains(normalize-space(.), {literal})])[1]",
            f"xpath=(//span[not(*) and contains(normalize-space(.), {literal})])[1]",
            f"xpath=(//p[not(*) and contains(normalize-space(.), {literal})])[1]",
            f"xpath=(//small[not(*) and contains(normalize-space(.), {literal})])[1]",
        ]
        for candidate in candidates:
            locator = page.locator(candidate).first
            if locator.count() > 0:
                return locator
        return None

    def _locator_is_interactable(self, locator, kind: str) -> bool:
        if locator is None:
            return False
        try:
            if locator.count() == 0:
                return False
        except Exception:
            return False

        if kind == "file":
            return True

        try:
            tag_name = str(locator.evaluate("(el) => (el.tagName || '').toLowerCase()") or "").lower()
        except Exception:
            tag_name = ""
        try:
            input_type = str(locator.evaluate("(el) => (el.getAttribute('type') || '').toLowerCase()") or "").lower()
        except Exception:
            input_type = ""
        try:
            role_name = str(locator.evaluate("(el) => (el.getAttribute('role') || '').toLowerCase()") or "").lower()
        except Exception:
            role_name = ""
        try:
            contenteditable = str(locator.evaluate("(el) => (el.getAttribute('contenteditable') || '').toLowerCase()") or "").lower()
        except Exception:
            contenteditable = ""
        try:
            disabled = bool(locator.evaluate("(el) => !!el.disabled || el.getAttribute('aria-disabled') === 'true'"))
        except Exception:
            disabled = False
        try:
            visible = bool(locator.is_visible())
        except Exception:
            visible = False

        if disabled or input_type == "hidden":
            return False
        if tag_name == "select":
            return True
        if kind == "editor":
            return visible or tag_name == "textarea" or contenteditable == "true" or role_name == "textbox"
        if kind in {"text", "dropdown"}:
            return visible or tag_name in {"input", "textarea"} or role_name in {"combobox", "textbox"}
        return visible or tag_name in {"input", "textarea"}

    def _resolve_control_locator(self, page, field_name: str, *, dynamic_text: str = ""):
        conf = FORM_SELECTORS[field_name]
        kind = conf.get("kind", "")
        for index, strategy in enumerate(conf.get("strategies", []), start=1):
            locator = self._locator_from_strategy(page, strategy, dynamic_text=dynamic_text)
            if locator is not None and self._locator_is_interactable(locator, kind):
                return locator, f"strategy#{index}:{strategy.get('type')}"

        label = conf.get("label", "")
        label_locator = self._find_label_locator(page, label)
        if label_locator is None or label_locator.count() == 0:
            return None, f"label_not_found:{label}"

        if kind in {"text", "dropdown"}:
            locator = label_locator.locator(
                "xpath=following::*[(self::input and not(@type='hidden') and not(@type='file')) "
                "or self::textarea or self::select or @role='combobox' or @role='textbox' or @contenteditable='true'][1]"
            ).first
            if self._locator_is_interactable(locator, kind):
                return locator, "label_following_input"
        if kind == "editor":
            editable = label_locator.locator("xpath=following::*[@contenteditable='true'][1]").first
            if self._locator_is_interactable(editable, kind):
                return editable, "label_following_editor"
            textarea = label_locator.locator("xpath=following::*[self::textarea][1]").first
            if self._locator_is_interactable(textarea, kind):
                return textarea, "label_following_textarea"
        if kind == "file":
            locator = label_locator.locator("xpath=following::*[self::input and @type='file'][1]").first
            if locator.count() > 0:
                return locator, "label_following_file"
        return None, f"locator_not_found:{field_name}"

    def _find_control_locator(self, page, field_name: str, *, dynamic_text: str = ""):
        locator, _ = self._resolve_control_locator(page, field_name, dynamic_text=dynamic_text)
        return locator

    def _is_login_page(self, page) -> bool:
        if "dang-nhap" in page.url.lower():
            return True
        username_locator = self._locator_from_strategy(page, LOGIN_SELECTORS["username"][0])
        password_locator = self._locator_from_strategy(page, LOGIN_SELECTORS["password"][0])
        try:
            return bool(
                username_locator
                and password_locator
                and username_locator.count() > 0
                and password_locator.count() > 0
            )
        except Exception:
            return False

    def ensure_authenticated(self, *, stop_event=None) -> None:
        self._ensure_context()
        page = self.anchor_page
        if self._page_is_closed(page):
            page = self.context.new_page()
        self.anchor_page = page
        page.goto(self.settings.create_url, wait_until="domcontentloaded")
        page_state = self._wait_for_create_or_login(page)
        if page_state == "login":
            self.log("[LOGIN] Session het han. Hay dang nhap lai tren Chromium.")
            if LOGIN_PATH not in str(page.url or "").lower():
                page.goto(self.settings.login_url, wait_until="domcontentloaded")
            self._begin_login_tracking(page)
            self.wait_for_manual_login(page, stop_event=stop_event)
            page.goto(self.settings.create_url, wait_until="domcontentloaded")
            page_state = self._wait_for_create_or_login(page)
        if page_state != "create":
            raise RuntimeError("Dang nhap xong nhung chua mo duoc trang tao ho so.")

    @property
    def staff_options_cache_path(self) -> Path:
        return self.working_dir / STAFF_OPTIONS_CACHE_NAME

    @staticmethod
    def load_staff_options_cache(working_dir: Path = BASE_DIR) -> dict[str, list[str]]:
        cache_path = Path(working_dir) / STAFF_OPTIONS_CACHE_NAME
        try:
            payload = json.loads(cache_path.read_text(encoding="utf-8"))
        except (OSError, ValueError, TypeError):
            return {"cong_chung_vien": [], "thu_ky": []}
        return {
            field_name: [str(label).strip() for label in payload.get(field_name, []) if str(label).strip()]
            for field_name in ("cong_chung_vien", "thu_ky")
        }

    def _read_dropdown_options(self, page, field_name: str) -> list[str]:
        locator, source = self._resolve_control_locator(page, field_name)
        if locator is None or locator.count() == 0:
            raise RuntimeError(f"Khong tim thay dropdown {field_name} ({source})")
        try:
            tag_name = str(locator.evaluate("(el) => el.tagName.toLowerCase()") or "").lower()
        except Exception:
            tag_name = ""
        if tag_name == "select":
            labels = locator.locator("option").all_text_contents()
        else:
            locator.click()
            page.wait_for_timeout(150)
            options = page.get_by_role("option")
            labels = options.all_text_contents() if options.count() else []
            try:
                page.keyboard.press("Escape")
            except Exception:
                pass

        result: list[str] = []
        for label in labels:
            text = " ".join(str(label or "").split())
            folded = _fold_value(text)
            if not text or folded in {"chon", "-- chon --", "chon cong chung vien", "chon thu ky"}:
                continue
            if text not in result:
                result.append(text)
        return result

    def fetch_staff_options(self) -> dict[str, list[str]]:
        self.ensure_authenticated()
        page = self.context.new_page()
        try:
            page.goto(self.settings.create_url, wait_until="domcontentloaded")
            self._wait_for_create_or_login(page)
            options = {
                field_name: self._read_dropdown_options(page, field_name)
                for field_name in ("cong_chung_vien", "thu_ky")
            }
            self.staff_options_cache_path.write_text(
                json.dumps(options, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            self.log(
                f"[UPLOAD] Da cap nhat danh sach CCV={len(options['cong_chung_vien'])}, thu ky={len(options['thu_ky'])}."
            )
            return options
        finally:
            page.close()

    def download_contract_book_export(
        self,
        *,
        from_date: str,
        to_date: str,
        download_dir: Path | None = None,
    ) -> Path:
        from_date_value = _parse_export_date_input(from_date)
        to_date_value = _parse_export_date_input(to_date)
        target_dir = Path(download_dir or (self.working_dir / DOWNLOADS_DIR.name))
        target_dir.mkdir(parents=True, exist_ok=True)

        self.ensure_authenticated()
        page = self.context.new_page()
        try:
            listing_url = f"{self.settings.base_url.rstrip('/')}/ho-so-cong-chung?page=1"
            page.goto(listing_url, wait_until="domcontentloaded")
            export_button = page.get_by_text("Xuất Sổ công chứng", exact=False).first
            deadline = time.monotonic() + 30
            while export_button.count() == 0 and time.monotonic() <= deadline:
                if self._is_login_page(page):
                    raise RuntimeError("Session da het han. Hay mo Cau hinh va dang nhap lai.")
                page.wait_for_timeout(250)
            if export_button.count() == 0:
                raise RuntimeError("Khong tim thay nut 'Xuat So cong chung' tren trang danh sach.")

            export_button.click()
            dialog = page.get_by_role("dialog").first
            if dialog.count() == 0:
                raise RuntimeError("Khong mo duoc popup chon thoi gian xuat so cong chung.")

            date_inputs = dialog.locator("input[placeholder='dd/mm/yyyy']")
            if date_inputs.count() < 2:
                raise RuntimeError("Khong tim thay du 2 o ngay bat dau/ket thuc trong popup export.")

            date_inputs.nth(0).click()
            date_inputs.nth(0).fill(from_date_value)
            date_inputs.nth(0).press("Tab")
            date_inputs.nth(1).click()
            date_inputs.nth(1).fill(to_date_value)
            date_inputs.nth(1).press("Tab")
            page.wait_for_timeout(300)

            download_button = dialog.get_by_role("button", name="Tải xuống", exact=False).first
            if download_button.count() == 0:
                raise RuntimeError("Khong tim thay nut 'Tải xuống' trong popup export.")

            with page.expect_download(timeout=60000) as download_info:
                download_button.click()
            download = download_info.value
            suggested_name = download.suggested_filename or f"so_cong_chung_{from_date_value.replace('/', '-')}_{to_date_value.replace('/', '-')}.xlsx"
            save_path = target_dir / suggested_name
            if save_path.exists():
                stem = save_path.stem
                suffix = save_path.suffix
                save_path = target_dir / f"{stem}_{now_iso().replace(':', '').replace('-', '')}{suffix}"
            download.save_as(str(save_path))
            self.log(
                f"[UPLOAD][EXPORT] Tai so cong chung tu web OK: {save_path} | from={from_date_value} | to={to_date_value}"
            )
            return save_path
        finally:
            page.close()

    def _fill_text(self, page, field_name: str, value: str) -> bool:
        locator, source = self._resolve_control_locator(page, field_name)
        if locator is None or locator.count() == 0:
            self.log(f"[UPLOAD][FIELD] {field_name}: locator not found ({source})")
            return False
        try:
            locator.click()
            locator.fill(value)
            locator.press("Tab")
            self.log(f"[UPLOAD][FIELD] {field_name}: filled via {source}")
            return True
        except Exception as exc:
            self.log(f"[UPLOAD][FIELD] {field_name}: fill_text failed via {source}: {exc}")
            return False

    def _dispatch_control_events(self, locator) -> None:
        try:
            locator.evaluate(
                """(el) => {
                    const hasValue = "value" in el;
                    if (hasValue) {
                        const proto = el instanceof HTMLTextAreaElement
                            ? HTMLTextAreaElement.prototype
                            : HTMLInputElement.prototype;
                        const descriptor = Object.getOwnPropertyDescriptor(proto, "value");
                        if (descriptor && descriptor.set) {
                            descriptor.set.call(el, el.value);
                        }
                    }
                    const inputEvent = typeof InputEvent === "function"
                        ? new InputEvent("input", {
                            bubbles: true,
                            inputType: "insertText",
                            data: hasValue ? el.value : "",
                        })
                        : new Event("input", { bubbles: true });
                    el.dispatchEvent(inputEvent);
                    el.dispatchEvent(new Event("change", { bubbles: true }));
                    el.dispatchEvent(new Event("blur", { bubbles: true }));
                }"""
            )
        except Exception:
            pass

    def _commit_control_events(self, page, locator) -> None:
        self._dispatch_control_events(locator)
        try:
            locator.press("Tab")
        except Exception:
            try:
                page.keyboard.press("Tab")
            except Exception:
                pass
        try:
            page.wait_for_timeout(200)
        except Exception:
            pass

    def _clear_control_like_user(self, page, locator) -> None:
        try:
            locator.click()
            locator.press("Control+A")
            locator.press("Backspace")
        except Exception:
            try:
                locator.fill("")
            except Exception:
                pass
        self._dispatch_control_events(locator)
        try:
            page.wait_for_timeout(100)
        except Exception:
            pass

    def _select_dropdown_option(self, page, locator, candidate: str) -> None:
        option = page.get_by_role("option", name=candidate, exact=False).first
        if option.count() == 0:
            option = page.get_by_text(candidate, exact=False).first
        if option.count() > 0:
            option.click()
            return
        try:
            locator.press("ArrowDown")
            page.wait_for_timeout(100)
        except Exception:
            pass
        locator.press("Enter")

    def _safe_field_validation_error(self, page, field_name: str) -> str:
        try:
            return self._read_field_validation_error(page, field_name).strip()
        except Exception:
            return ""

    def _dropdown_acceptance_state(self, page, field_name: str, expected: str) -> tuple[bool, str, str]:
        actual = self._read_field_value(page, field_name).strip()
        validation_error = self._safe_field_validation_error(page, field_name)
        accepted = field_value_matches(field_name, expected, actual) and not validation_error
        return accepted, actual, validation_error

    def _prime_contract_name_validation(self, page) -> bool:
        try:
            if self._read_field_value(page, "so_cong_chung") or self._read_field_value(page, "nhom_hop_dong"):
                return False
        except Exception:
            return False

        for strategy in SAVE_BUTTON_SELECTORS:
            button = self._locator_from_strategy(page, strategy)
            try:
                if button is not None and button.count() > 0 and button.is_visible():
                    button.click(timeout=5000)
                    page.wait_for_timeout(500)
                    return True
            except Exception:
                continue
        return False

    def _type_dropdown_candidate(self, page, locator, candidate: str) -> None:
        locator.click()
        self._clear_control_like_user(page, locator)
        locator.type(candidate, delay=5)
        page.wait_for_timeout(150)
        self._select_dropdown_option(page, locator, candidate)
        self._commit_control_events(page, locator)

    def _retype_contract_name_after_validation(self, page, locator, candidate: str) -> None:
        locator.click()
        locator.press("Control+A")
        locator.press("Backspace")
        page.wait_for_timeout(100)
        locator.type(candidate, delay=5)
        locator.press("Tab")
        page.wait_for_timeout(500)

    def _fill_dropdown(self, page, field_name: str, value: str) -> bool:
        locator, source = self._resolve_control_locator(page, field_name)
        if locator is None or locator.count() == 0:
            self.log(f"[UPLOAD][FIELD] {field_name}: locator not found ({source})")
            return False
        candidates = get_field_value_candidates(field_name, value)
        if not candidates:
            self.log(f"[UPLOAD][FIELD] {field_name}: empty dropdown value")
            return False

        try:
            tag_name = locator.evaluate("(el) => el.tagName.toLowerCase()")
        except Exception:
            tag_name = ""
        if str(tag_name).lower() == "select":
            last_error = ""
            for candidate in candidates:
                try:
                    locator.select_option(label=candidate)
                    page.wait_for_timeout(150)
                    self._commit_control_events(page, locator)
                    accepted, actual, validation_error = self._dropdown_acceptance_state(page, field_name, candidate)
                    if accepted:
                        self.log(f"[UPLOAD][FIELD] {field_name}: select filled via {source} using {candidate!r}")
                        return True
                    if validation_error:
                        last_error = (
                            f"select_validation expected={candidate!r} actual={actual!r} error={validation_error!r}"
                        )
                    else:
                        last_error = f"select_not_stuck expected={candidate!r} actual={actual!r}"
                except Exception as exc:
                    last_error = str(exc)
                    self.log(
                        f"[UPLOAD][FIELD] {field_name}: select_option failed via {source} with {candidate!r}: {exc}"
                    )
            self.log(f"[UPLOAD][FIELD] {field_name}: dropdown fill failed via {source}: {last_error}")
            return False

        current_text = self._read_field_value(page, field_name).strip()
        if tag_name != "input" and field_value_matches(field_name, value, current_text):
            current_validation_error = self._safe_field_validation_error(page, field_name)
            if not current_validation_error:
                self.log(f"[UPLOAD][FIELD] {field_name}: already set via {source}")
                return True

        last_error = ""
        for candidate in candidates:
            for attempt in range(2):
                try:
                    try:
                        page.keyboard.press("Escape")
                        page.wait_for_timeout(100)
                    except Exception:
                        pass

                    self._type_dropdown_candidate(page, locator, candidate)
                except Exception as exc:
                    try:
                        locator.click()
                        locator.fill(candidate)
                    except Exception as fallback_exc:
                        last_error = str(fallback_exc)
                        self.log(
                            f"[UPLOAD][FIELD] {field_name}: fill_dropdown typing failed via {source} with {candidate!r}: {exc}"
                        )
                        continue

                accepted, actual, validation_error = self._dropdown_acceptance_state(page, field_name, candidate)
                if accepted:
                    if field_name == "ten_hop_dong" and attempt == 0 and self._prime_contract_name_validation(page):
                        self._retype_contract_name_after_validation(page, locator, candidate)
                        accepted, actual, validation_error = self._dropdown_acceptance_state(
                            page,
                            field_name,
                            candidate,
                        )
                        if not accepted:
                            last_error = (
                                f"validation_after_prime expected={candidate!r} actual={actual!r} "
                                f"error={validation_error!r}"
                            )
                            continue
                    retry_text = " after retry" if attempt else ""
                    self.log(
                        f"[UPLOAD][FIELD] {field_name}: dropdown filled via {source} using {candidate!r}{retry_text}"
                    )
                    return True

                if validation_error:
                    last_error = f"validation expected={candidate!r} actual={actual!r} error={validation_error!r}"
                else:
                    last_error = f"value_not_stuck expected={candidate!r} actual={actual!r}"
                self.log(
                    f"[UPLOAD][FIELD] {field_name}: dropdown candidate {candidate!r} not accepted via {source}; actual={actual!r}; validation={validation_error!r}"
                )

        self.log(f"[UPLOAD][FIELD] {field_name}: dropdown fill failed via {source}: {last_error}")
        return False

    def _fill_editor(self, page, field_name: str, value: str) -> bool:
        locator, source = self._resolve_control_locator(page, field_name)
        if locator is None or locator.count() == 0:
            self.log(f"[UPLOAD][FIELD] {field_name}: locator not found ({source})")
            return False
        try:
            tag_name = locator.evaluate("(el) => el.tagName.toLowerCase()")
        except Exception:
            tag_name = ""

        if tag_name == "textarea":
            try:
                locator.fill(value)
                self.log(f"[UPLOAD][FIELD] {field_name}: textarea filled via {source}")
                return True
            except Exception as exc:
                self.log(f"[UPLOAD][FIELD] {field_name}: textarea fill failed via {source}: {exc}")
                return False

        try:
            locator.click()
            locator.fill(value)
            locator.press("Tab")
            self.log(f"[UPLOAD][FIELD] {field_name}: editor filled via {source}")
            return True
        except Exception as exc:
            self.log(f"[UPLOAD][FIELD] {field_name}: editor fill failed via {source}: {exc}")
            return False

    def _upload_file(self, page, file_path: Path) -> bool:
        locator, source = self._resolve_control_locator(page, "file_hop_dong", dynamic_text=file_path.name)
        if locator is None or locator.count() == 0:
            self.log(f"[UPLOAD][FIELD] file_hop_dong: locator not found ({source})")
            return False
        try:
            locator.set_input_files(str(file_path))
            page.wait_for_function(
                """(selector) => {
                    const el = document.querySelector(selector);
                    return !!(el && el.files && el.files.length > 0);
                }""",
                arg="#hopdong\\.fileHopdong",
                timeout=5000,
            )
            self.log(f"[UPLOAD][FIELD] file_hop_dong: uploaded via {source} -> {file_path.name}")
            return True
        except Exception as exc:
            success = page.get_by_text(file_path.name, exact=False).count() > 0
            self.log(
                f"[UPLOAD][FIELD] file_hop_dong: upload {'ok' if success else 'failed'} via {source}: {exc}"
            )
            return success

    def _read_field_value(self, page, field_name: str) -> str:
        if field_name == "file_hop_dong":
            locator, _ = self._resolve_control_locator(page, field_name)
            if locator is None or locator.count() == 0:
                return ""
            try:
                return locator.evaluate("(el) => (el.files && el.files.length > 0) ? el.files[0].name : ''")
            except Exception:
                return ""

        locator, _ = self._resolve_control_locator(page, field_name)
        if locator is None or locator.count() == 0:
            return ""
        try:
            tag_name = locator.evaluate("(el) => el.tagName.toLowerCase()")
        except Exception:
            tag_name = ""

        if tag_name == "select":
            try:
                return str(
                    locator.evaluate("(el) => el.selectedOptions.length ? el.selectedOptions[0].textContent : ''")
                    or ""
                ).strip()
            except Exception:
                return ""
        if tag_name in {"input", "textarea"}:
            try:
                return str(locator.input_value()).strip()
            except Exception:
                return ""
        try:
            return str(locator.text_content() or "").strip()
        except Exception:
            return ""

    def _read_field_validation_error(self, page, field_name: str) -> str:
        label = str(FORM_SELECTORS.get(field_name, {}).get("label") or "").strip()
        if not label:
            return ""
        label_locator = self._find_label_locator(page, label)
        if label_locator is None or label_locator.count() == 0:
            return ""

        scopes = [
            label_locator.locator("xpath=ancestor::*[contains(@class,'Field-container')][1]").first,
            label_locator.locator("xpath=ancestor::div[1]").first,
        ]
        seen: set[str] = set()
        for scope in scopes:
            try:
                if scope.count() == 0:
                    continue
                candidates = scope.locator("xpath=.//*[self::small or self::p][normalize-space()]")
                count = candidates.count()
            except Exception:
                continue
            for index in range(count):
                candidate = candidates.nth(index)
                try:
                    if not candidate.is_visible():
                        continue
                    text = candidate.text_content()
                except Exception:
                    continue
                normalized = _fold_value(text)
                if not normalized or normalized in seen:
                    continue
                seen.add(normalized)
                if "khong duoc de trong" in normalized or "bat buoc" in normalized:
                    return str(text or "").strip()
        return ""

    def _verify_record(self, page, record: UploadRecord) -> tuple[dict, bool]:
        verify_data = {"missing_fields": list(record.missing_fields), "fields": {}}
        partial = bool(record.missing_fields)
        for field_name in VERIFY_FIELDS:
            expected = record.upload_form.get(field_name, "")
            actual = self._read_field_value(page, field_name)
            if field_name == "file_hop_dong":
                success = bool(actual)
                expected_value = Path(expected).name if expected else ""
            else:
                expected_value = str(expected or "").strip()
                if not expected_value:
                    success = False
                elif FORM_SELECTORS.get(field_name, {}).get("kind") == "dropdown":
                    success = field_value_matches(field_name, expected_value, actual)
                else:
                    success = expected_value.lower() in str(actual or "").lower()
            validation_error = self._read_field_validation_error(page, field_name)
            if validation_error:
                success = False
            verify_data["fields"][field_name] = {
                "expected": expected_value,
                "actual": actual,
                "success": success,
                "validation_error": validation_error,
            }
            self.log(
                f"[UPLOAD][VERIFY] {record.contract_no} {field_name}: "
                f"expected={expected_value!r} actual={actual!r} success={success}"
                + (f" validation_error={validation_error!r}" if validation_error else "")
            )
            if not success:
                partial = True
        return verify_data, partial

    @property
    def open_prepared_record_ids(self) -> set[int]:
        return set(self._prepared_record_ids)

    def _register_prepared_page(
        self,
        record: UploadRecord,
        page,
        artifact_dir: Path,
    ) -> None:
        record_id = int(record.record_id)
        network_log_path = artifact_dir / "network.jsonl"
        tab = PreparedBrowserTab(
            record_id=record_id,
            contract_no=str(record.contract_no),
            page=page,
            initial_url=str(page.url or self.settings.create_url),
            network_log_path=network_log_path,
        )
        self.prepared_pages[record_id] = tab
        self._prepared_record_ids.add(record_id)

        def handle_response(response) -> None:
            try:
                request = response.request
                method = str(request.method or "").upper()
                url = str(response.url or "")
                status = int(response.status or 0)
                path = urllib.parse.urlsplit(url).path.rstrip("/")
                if method == "POST" and path == SAVE_API_PATH:
                    self._append_network_event(
                        network_log_path,
                        method=method,
                        url=url,
                        status=status,
                        record_id=record_id,
                    )
                    if 200 <= status < 300:
                        tab.save_response_ok = True
            except Exception:
                pass

        try:
            page.on("response", handle_response)
        except Exception:
            pass

    def _page_left_create_route(self, tab: PreparedBrowserTab) -> bool:
        current_url = str(getattr(tab.page, "url", "") or "")
        current = urllib.parse.urlsplit(current_url)
        base = urllib.parse.urlsplit(self.settings.base_url)
        current_path = current.path.rstrip("/")
        create_path = urllib.parse.urlsplit(self.settings.create_url).path.rstrip("/")
        return bool(
            current.scheme in {"http", "https"}
            and current.netloc == base.netloc
            and current_path != create_path
            and current_path != LOGIN_PATH
        )

    def poll_prepared_pages(self) -> dict[str, list[int]]:
        saved_ids: list[int] = []
        closed_ids: list[int] = []
        for record_id, tab in list(self.prepared_pages.items()):
            page = tab.page
            if self._page_is_closed(page):
                self.prepared_pages.pop(record_id, None)
                self._prepared_record_ids.discard(record_id)
                closed_ids.append(record_id)
                self.log(
                    f"[UPLOAD] Tab {tab.contract_no} da dong khi chua nhan duoc tin hieu Luu; "
                    "ho so co the duoc chuan bi lai."
                )
                continue

            try:
                # A small Playwright call lets the sync API dispatch response events.
                page.evaluate("() => document.readyState")
            except Exception:
                continue

            if not tab.save_response_ok and not self._page_left_create_route(tab):
                continue

            finalize_uploaded_records([record_id], working_dir=self.working_dir)
            self.prepared_pages.pop(record_id, None)
            self._prepared_record_ids.discard(record_id)
            saved_ids.append(record_id)
            signal = "POST /api/hoso 2xx" if tab.save_response_ok else "trang da chuyen sau khi Luu"
            self.log(f"[UPLOAD] Da nhan dien Luu {tab.contract_no} ({signal}).")
            try:
                page.close()
            except Exception:
                pass

        return {
            "saved_record_ids": saved_ids,
            "closed_record_ids": closed_ids,
            "open_record_ids": sorted(self._prepared_record_ids),
        }

    def _prepare_record(self, record: UploadRecord, artifact_dir: Path):
        if not record.source_file.exists():
            self.log(f"[UPLOAD] File goc khong ton tai, tiep tuc khong upload file: {record.source_file}")

        page = self.context.new_page()
        page.goto(self.settings.create_url, wait_until="domcontentloaded")
        page_state = self._wait_for_create_or_login(page)
        if page_state == "login":
            self._begin_login_tracking(page)
            self.wait_for_manual_login(page)
            page.goto(self.settings.create_url, wait_until="domcontentloaded")
            page_state = self._wait_for_create_or_login(page)
        if page_state != "create":
            raise RuntimeError("Khong mo duoc form tao ho so sau khi dang nhap.")

        field_results: dict[str, dict] = {}
        for field_name in FORM_FIELD_ORDER:
            value = str(record.upload_form.get(field_name, "") or "")
            if field_name == "file_hop_dong":
                if value:
                    field_results[field_name] = {"value": Path(value).name, "success": self._upload_file(page, Path(value))}
                continue
            if not value:
                field_results[field_name] = {"value": "", "success": False, "reason": "empty_value"}
                continue
            field_kind = FORM_SELECTORS[field_name]["kind"]
            if field_kind == "text":
                field_results[field_name] = {"value": value, "success": self._fill_text(page, field_name, value)}
            elif field_kind == "dropdown":
                field_results[field_name] = {"value": value, "success": self._fill_dropdown(page, field_name, value)}
            elif field_kind == "editor":
                field_results[field_name] = {"value": value, "success": self._fill_editor(page, field_name, value)}

        verify_data, partial = self._verify_record(page, record)
        screenshot_path = artifact_dir / f"before_save_{sanitize_contract_no(record.contract_no)}.png"
        debug_json_path = artifact_dir / f"debug_{sanitize_contract_no(record.contract_no)}.json"
        page.screenshot(path=str(screenshot_path), full_page=True)
        debug_payload = {
            "contract_no": record.contract_no,
            "source_file": str(record.source_file),
            "page_url": page.url,
            "field_results": field_results,
            "verify": verify_data,
        }
        debug_json_path.write_text(json.dumps(debug_payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if self.settings.post_prepare_delay_ms > 0:
            page.wait_for_timeout(self.settings.post_prepare_delay_ms)

        if getattr(record, "record_id", None) is not None:
            self._register_prepared_page(record, page, artifact_dir)

        return {
            "status": "prepared_partial" if partial else "prepared_dry_run",
            "verify_json": json.dumps(verify_data, ensure_ascii=False, indent=2),
            "artifact_dir": str(artifact_dir),
            "screenshot": str(screenshot_path),
            "debug_json": str(debug_json_path),
        }

    def prepare_manifest(
        self,
        manifest_path: Path | str,
        stop_event,
        *,
        selected_record_ids: set[int] | None = None,
        exclude_contract_nos: Optional[set[str]] = None,
        requester_sheet_url: str = "",
        progress_callback: Optional[Callable[[dict], None]] = None,
        cong_chung_vien: str | None = None,
        thu_ky: str | None = None,
        chunk_size: int | None = None,
    ) -> dict:
        def emit_progress(event: str, **payload: object) -> None:
            if progress_callback is None:
                return
            progress_payload = {"event": event, **payload}
            try:
                progress_callback(progress_payload)
            except Exception:
                # UI progress updates must not break the uploader flow.
                pass

        requester_lookup, requester_warning = load_requester_contract_lookup(
            working_dir=self.working_dir,
            sheet_url=requester_sheet_url,
            log_callback=self.log,
        )
        if requester_warning:
            self.log(f"[UPLOAD][REQUESTER] {requester_warning}")

        manifest, records, total_pending = load_upload_queue(
            manifest_path,
            working_dir=self.working_dir,
            selected_record_ids=selected_record_ids,
            requester_lookup=requester_lookup,
            cong_chung_vien=cong_chung_vien,
            thu_ky=thu_ky,
            statuses=QUEUE_STATUSES,
        )
        run_id = str(manifest["run_id"])
        records = [
            record
            for record in records
            if record.status in PREPARE_QUEUE_STATUSES
            or (
                record.status in PREPARED_QUEUE_STATUSES
                and record.record_id not in self._prepared_record_ids
            )
        ]
        filtered_records, duplicate_records = split_records_by_existing_contract_nos(
            records,
            set(exclude_contract_nos or set()),
        )
        effective_chunk_size = max(1, min(int(chunk_size or self.settings.max_prepared_tabs), 30))
        records = filtered_records[:effective_chunk_size]
        filtered_pending = len(filtered_records)
        emit_progress(
            "queue_loaded",
            run_id=run_id,
            total_pending=total_pending,
            filtered_pending=filtered_pending,
            chunk_size=len(records),
            excluded_duplicates=len(duplicate_records),
            manifest_path=str(Path(manifest_path)),
        )

        if not records:
            duplicate_message = ""
            if duplicate_records:
                duplicate_message = f" Da loai {len(duplicate_records)} ho so trung so cong chung tren web."
            summary = {
                "run_id": run_id,
                "prepared_count": 0,
                "total_pending": total_pending,
                "remaining": 0,
                "artifact_dir": "",
                "excluded_duplicates": len(duplicate_records),
                "open_record_ids": sorted(self._prepared_record_ids),
                "message": f"Khong con ho so nao can chuan bi.{duplicate_message}",
            }
            emit_progress("finished", **summary)
            return summary

        artifact_dir = self.working_dir / "upload_runs" / f"{now_iso().replace(':', '').replace('-', '')}_{run_id}"
        artifact_dir.mkdir(parents=True, exist_ok=True)
        self.run_log_path = artifact_dir / "dry_run_trace.log"
        self.log(f"[UPLOAD] Artifact dir: {artifact_dir}")
        self.ensure_authenticated(stop_event=stop_event)

        prepared_count = 0
        errors = []
        conn = connect_registry(self.working_dir / REGISTRY_DB_PATH.name)
        try:
            for record in records:
                if stop_event.is_set():
                    emit_progress(
                        "stopped",
                        run_id=run_id,
                        prepared_count=prepared_count,
                        total_pending=filtered_pending,
                        remaining=max(filtered_pending - prepared_count, 0),
                        artifact_dir=str(artifact_dir),
                        excluded_duplicates=len(duplicate_records),
                    )
                    break
                emit_progress(
                    "record_started",
                    run_id=run_id,
                    record_id=record.record_id,
                    contract_no=record.contract_no,
                    prepared_count=prepared_count,
                    total_pending=filtered_pending,
                    remaining=max(filtered_pending - prepared_count, 0),
                    artifact_dir=str(artifact_dir),
                )
                self.log(f"[UPLOAD] Dang chuan bi {record.contract_no} (record #{record.record_id})")
                try:
                    result = self._prepare_record(record, artifact_dir)
                    self._prepared_record_ids.add(record.record_id)
                    update_registry_record_by_id(
                        conn,
                        record.record_id,
                        status=result["status"],
                        reason="Dry-run prepared",
                        last_error="",
                        prepared_at=now_iso(),
                        artifact_dir=result["artifact_dir"],
                        verify_json=result["verify_json"],
                    )
                    prepared_count += 1
                    self.log(
                        f"[UPLOAD] {record.contract_no}: {result['status']} -> {result['screenshot']} | debug={result['debug_json']}"
                    )
                    emit_progress(
                        "record_prepared",
                        run_id=run_id,
                        record_id=record.record_id,
                        contract_no=record.contract_no,
                        status=result["status"],
                        prepared_count=prepared_count,
                        total_pending=filtered_pending,
                        remaining=max(filtered_pending - prepared_count, 0),
                        artifact_dir=str(artifact_dir),
                        screenshot=result["screenshot"],
                        debug_json=result["debug_json"],
                    )
                except Exception as exc:
                    errors.append({"record_id": record.record_id, "contract_no": record.contract_no, "error": str(exc)})
                    update_registry_record_by_id(
                        conn,
                        record.record_id,
                        status="upload_failed",
                        last_error=str(exc),
                        reason="Uploader dry-run failed",
                        artifact_dir=str(artifact_dir),
                    )
                    self.log(f"[UPLOAD] Loi {record.contract_no}: {exc}")
                    emit_progress(
                        "record_failed",
                        run_id=run_id,
                        record_id=record.record_id,
                        contract_no=record.contract_no,
                        error=str(exc),
                        prepared_count=prepared_count,
                        total_pending=filtered_pending,
                        remaining=max(filtered_pending - prepared_count, 0),
                        artifact_dir=str(artifact_dir),
                    )
            summary = {
                "run_id": run_id,
                "prepared_count": prepared_count,
                "total_pending": filtered_pending,
                "remaining": max(filtered_pending - prepared_count, 0),
                "artifact_dir": str(artifact_dir),
                "excluded_duplicates": len(duplicate_records),
                "open_record_ids": sorted(self._prepared_record_ids),
                "errors": errors,
                "message": (
                    f"Da chuan bi {prepared_count}/{filtered_pending} ho so."
                    + (f" Da loai {len(duplicate_records)} ho so trung tren web." if duplicate_records else "")
                    + " Hay ra soat va bam Luu tren web; app se tu cap nhat."
                ),
            }
            (artifact_dir / "upload_manifest.json").write_text(
                json.dumps(summary, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            emit_progress("finished", **summary)
            return summary
        finally:
            self.run_log_path = None
            conn.close()


def download_contract_book_export(
    *,
    from_date: str,
    to_date: str,
    working_dir: Path = BASE_DIR,
    settings: Optional[UploaderSettings] = None,
    log_callback: Optional[Callable[[str], None]] = None,
) -> Path:
    session = NamDinhUploaderSession(
        settings or load_uploader_settings(working_dir),
        working_dir=working_dir,
        log_callback=log_callback,
    )
    try:
        return session.download_contract_book_export(from_date=from_date, to_date=to_date)
    finally:
        session.close()
