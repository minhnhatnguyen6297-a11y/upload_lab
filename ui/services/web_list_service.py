from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from playwright_uploader import normalize_contract_no_for_compare


@dataclass(frozen=True)
class ContractListRow:
    row_index: int
    raw_value: str
    contract_no: str
    year: int
    ordinal: int


@dataclass(frozen=True)
class MissingContractNo:
    year: int
    ordinal: int
    contract_no: str


@dataclass(frozen=True)
class ContractLookupResult:
    found: bool
    query: str
    contract_no: str = ""
    row_index: int | None = None
    raw_value: str = ""
    message: str = ""


def _header_key(value: str) -> str:
    return re.sub(r"[^A-Z]", "", value.upper())


def _parse_normalized_contract_no(value: str) -> tuple[int, int] | None:
    match = re.fullmatch(r"(\d+)/(\d{4})", value)
    if not match:
        return None
    return int(match.group(2)), int(match.group(1))


def read_exported_contract_rows(export_path: Path | str) -> list[ContractListRow]:
    path = Path(export_path)
    if not path.exists():
        raise FileNotFoundError(f"Khong tim thay file Excel: {path}")

    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - dependency wiring
        raise RuntimeError("Thieu openpyxl trong moi truong tool.") from exc

    rows: list[ContractListRow] = []
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        for row_index, (cell_value,) in enumerate(
            sheet.iter_rows(min_col=1, max_col=1, values_only=True),
            start=1,
        ):
            raw_value = str(cell_value or "").strip()
            if not raw_value:
                continue
            if row_index == 1 and "SOCONGCHUNG" in _header_key(raw_value):
                continue
            contract_no = normalize_contract_no_for_compare(raw_value)
            parsed = _parse_normalized_contract_no(contract_no)
            if not parsed:
                continue
            year, ordinal = parsed
            rows.append(
                ContractListRow(
                    row_index=row_index,
                    raw_value=raw_value,
                    contract_no=contract_no,
                    year=year,
                    ordinal=ordinal,
                )
            )
    finally:
        workbook.close()
    return rows


def find_missing_contract_numbers(rows: list[ContractListRow]) -> list[MissingContractNo]:
    ordinals_by_year: dict[int, set[int]] = {}
    for row in rows:
        ordinals_by_year.setdefault(row.year, set()).add(row.ordinal)

    missing: list[MissingContractNo] = []
    for year in sorted(ordinals_by_year):
        ordinals = ordinals_by_year[year]
        if len(ordinals) < 2:
            continue
        for ordinal in range(min(ordinals), max(ordinals) + 1):
            if ordinal not in ordinals:
                missing.append(
                    MissingContractNo(
                        year=year,
                        ordinal=ordinal,
                        contract_no=f"{ordinal}/{year}",
                    )
                )
    return missing


def lookup_exported_contract_no(rows: list[ContractListRow], query: str) -> ContractLookupResult:
    normalized = normalize_contract_no_for_compare(query)
    if not normalized:
        return ContractLookupResult(found=False, query=query, message="So cong chung khong hop le.")

    for row in rows:
        if row.contract_no == normalized:
            return ContractLookupResult(
                found=True,
                query=query,
                contract_no=row.contract_no,
                row_index=row.row_index,
                raw_value=row.raw_value,
                message=f"Tim thay dong {row.row_index}.",
            )
    return ContractLookupResult(found=False, query=query, contract_no=normalized, message="Khong tim thay.")
