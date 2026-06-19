from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from enum import Enum
from pathlib import Path
import re
import unicodedata


class ContractBookIssueKind(str, Enum):
    BAD_FORMAT = "sai_format"
    WRONG_YEAR = "sai_nam"
    DUPLICATE = "trung_so"
    DATE_PARSE_ERROR = "ngay_khong_hop_le"
    DATE_SEQUENCE = "ngay_bat_thuong"
    SAME_DAY_JUMP = "nhay_bat_thuong_cung_ngay"


ContractBookWarningKind = ContractBookIssueKind


@dataclass(frozen=True)
class ContractBookDisplayRow:
    row_index: int
    raw_contract_no: str
    raw_date: str
    contract_no: str
    year: int
    ordinal: int
    contract_date: date


ContractBookRow = ContractBookDisplayRow


@dataclass(frozen=True)
class MissingContractNumber:
    year: int
    ordinal: int
    contract_no: str


@dataclass(frozen=True)
class ContractBookIssueRow:
    row_index: int
    kind: ContractBookIssueKind
    raw_contract_no: str
    raw_date: str
    contract_no: str
    year: int | None
    ordinal: int | None
    contract_date: date | None
    message: str


ContractBookWarning = ContractBookIssueRow


@dataclass(frozen=True)
class ContractBookSummary:
    excel_total: int
    valid_count: int
    missing_count: int
    issue_count: int
    duplicate_count: int


@dataclass(frozen=True)
class ContractBookAnalysis:
    source_path: Path
    display_rows: list[ContractBookDisplayRow]
    missing_numbers: list[MissingContractNumber]
    issue_rows: list[ContractBookIssueRow]
    summary: ContractBookSummary
    min_ordinal: int | None
    max_ordinal: int | None

    @property
    def valid_rows(self) -> list[ContractBookDisplayRow]:
        return self.display_rows

    @property
    def warning_rows(self) -> list[ContractBookIssueRow]:
        return self.issue_rows

    @property
    def parse_error_rows(self) -> list[ContractBookIssueRow]:
        return [issue for issue in self.issue_rows if issue.kind == ContractBookIssueKind.BAD_FORMAT]


def _fold_text(value: object) -> str:
    normalized = unicodedata.normalize("NFD", str(value or ""))
    normalized = normalized.replace("\u0111", "d").replace("\u0110", "D")
    folded = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
    return folded.upper()


def _header_key(value: object) -> str:
    return re.sub(r"[^A-Z]", "", _fold_text(value))


def _stringify_cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_contract_book_no(value: object, *, default_year: int | None = None) -> tuple[int, int, str] | None:
    del default_year
    text = re.sub(r"\s+", "", _fold_text(value))
    if not text:
        return None
    match = re.fullmatch(r"0*(\d+)/(\d{4})(?:/CCGD)?", text)
    if not match:
        return None
    ordinal = int(match.group(1))
    year = int(match.group(2))
    return year, ordinal, f"{ordinal}/{year}"


def parse_contract_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value or "").strip()
    if not text:
        return None

    for pattern in ("%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _parse_required_date(value: str | date, label: str) -> date:
    parsed = parse_contract_date(value)
    if parsed is None:
        raise ValueError(f"{label} khong dung dinh dang dd/mm/yyyy: {value}")
    return parsed


def _allowed_years(from_date: date, to_date: date) -> set[int]:
    if to_date < from_date:
        raise ValueError("Ngay ket thuc nho hon ngay bat dau.")
    return set(range(from_date.year, to_date.year + 1))


def _read_rows(path: Path) -> list[tuple[int, object, object]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - dependency wiring
        raise RuntimeError("Thieu openpyxl trong moi truong tool.") from exc

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(min_col=1, max_col=2, values_only=True))
    finally:
        workbook.close()

    start_index = 1
    if rows:
        header_a = _header_key(rows[0][0])
        if "SOCONGCHUNG" in header_a:
            start_index = 2

    return [
        (row_index, row[0], row[1])
        for row_index, row in enumerate(rows[start_index - 1 :], start=start_index)
    ]


def _build_missing_numbers(rows: list[ContractBookDisplayRow]) -> list[MissingContractNumber]:
    ordinals_by_year: dict[int, set[int]] = {}
    for row in rows:
        ordinals_by_year.setdefault(row.year, set()).add(row.ordinal)

    missing: list[MissingContractNumber] = []
    for year in sorted(ordinals_by_year):
        ordinals = ordinals_by_year[year]
        if len(ordinals) < 2:
            continue
        for ordinal in range(min(ordinals), max(ordinals) + 1):
            if ordinal not in ordinals:
                missing.append(
                    MissingContractNumber(
                        year=year,
                        ordinal=ordinal,
                        contract_no=f"{ordinal}/{year}",
                    )
                )
    return missing


def _filter_missing_numbers(
    missing_numbers: list[MissingContractNumber],
    issues: list[ContractBookIssueRow],
) -> list[MissingContractNumber]:
    blocked_ordinals = {
        (int(issue.year), int(issue.ordinal))
        for issue in issues
        if issue.year is not None and issue.ordinal is not None
    }
    return [
        item
        for item in missing_numbers
        if (item.year, item.ordinal) not in blocked_ordinals
    ]


def _make_issue(
    row: ContractBookDisplayRow,
    kind: ContractBookIssueKind,
    message: str,
) -> ContractBookIssueRow:
    return ContractBookIssueRow(
        row_index=row.row_index,
        kind=kind,
        raw_contract_no=row.raw_contract_no,
        raw_date=row.raw_date,
        contract_no=row.contract_no,
        year=row.year,
        ordinal=row.ordinal,
        contract_date=row.contract_date,
        message=message,
    )


def _remove_date_sequence_issues(
    rows: list[ContractBookDisplayRow],
    issues: list[ContractBookIssueRow],
) -> list[ContractBookDisplayRow]:
    clean: list[ContractBookDisplayRow] = []
    latest_date: date | None = None
    for row in sorted(rows, key=lambda item: (item.year, item.ordinal, item.row_index)):
        if latest_date is not None and row.contract_date < latest_date:
            issues.append(
                _make_issue(
                    row,
                    ContractBookIssueKind.DATE_SEQUENCE,
                    "Ngay cong chung som hon ngay cua so nho hon truoc do.",
                )
            )
            continue
        latest_date = row.contract_date if latest_date is None else max(latest_date, row.contract_date)
        clean.append(row)
    return clean


def _remove_same_day_jump_issues(
    rows: list[ContractBookDisplayRow],
    issues: list[ContractBookIssueRow],
) -> list[ContractBookDisplayRow]:
    issue_ids: set[tuple[int, str]] = set()
    rows_by_date: dict[date, list[ContractBookDisplayRow]] = {}
    for row in rows:
        rows_by_date.setdefault(row.contract_date, []).append(row)

    for contract_date in sorted(rows_by_date):
        previous: ContractBookDisplayRow | None = None
        for row in sorted(rows_by_date[contract_date], key=lambda item: (item.ordinal, item.row_index)):
            if previous is not None and row.ordinal - previous.ordinal > 10:
                issues.append(
                    _make_issue(
                        row,
                        ContractBookIssueKind.SAME_DAY_JUMP,
                        f"So nhay {row.ordinal - previous.ordinal} trong cung ngay, vuot nguong 10.",
                    )
                )
                issue_ids.add((row.row_index, row.contract_no))
                continue
            previous = row

    return [row for row in rows if (row.row_index, row.contract_no) not in issue_ids]


def analyze_contract_book(
    export_path: Path | str,
    *,
    from_date: str | date,
    to_date: str | date,
) -> ContractBookAnalysis:
    path = Path(export_path)
    if not path.exists():
        raise FileNotFoundError(f"Khong tim thay file Excel: {path}")

    start_date = _parse_required_date(from_date, "from_date")
    end_date = _parse_required_date(to_date, "to_date")
    years = _allowed_years(start_date, end_date)

    parsed_rows: list[ContractBookDisplayRow] = []
    issues: list[ContractBookIssueRow] = []
    excel_total = 0

    for row_index, raw_contract_no, raw_date in _read_rows(path):
        raw_contract_text = _stringify_cell_value(raw_contract_no)
        raw_date_text = _stringify_cell_value(raw_date)
        if not raw_contract_text and not raw_date_text:
            continue
        excel_total += 1

        parsed_no = parse_contract_book_no(raw_contract_no)
        contract_date = parse_contract_date(raw_date)
        if parsed_no is None:
            issues.append(
                ContractBookIssueRow(
                    row_index=row_index,
                    kind=ContractBookIssueKind.BAD_FORMAT,
                    raw_contract_no=raw_contract_text,
                    raw_date=raw_date_text,
                    contract_no="",
                    year=None,
                    ordinal=None,
                    contract_date=contract_date,
                    message="So cong chung khong dung dang xxx/yyyy hoac xxx/yyyy/CCGD.",
                )
            )
            continue

        year, ordinal, contract_no = parsed_no
        if contract_date is None:
            issues.append(
                ContractBookIssueRow(
                    row_index=row_index,
                    kind=ContractBookIssueKind.DATE_PARSE_ERROR,
                    raw_contract_no=raw_contract_text,
                    raw_date=raw_date_text,
                    contract_no=contract_no,
                    year=year,
                    ordinal=ordinal,
                    contract_date=None,
                    message="Ngay cot B khong dung dinh dang dd/mm/yyyy.",
                )
            )
            continue
        if year not in years:
            issues.append(
                ContractBookIssueRow(
                    row_index=row_index,
                    kind=ContractBookIssueKind.WRONG_YEAR,
                    raw_contract_no=raw_contract_text,
                    raw_date=raw_date_text,
                    contract_no=contract_no,
                    year=year,
                    ordinal=ordinal,
                    contract_date=contract_date,
                    message=f"Nam {year} khong thuoc khoang ngay dang chon.",
                )
            )
            continue

        parsed_rows.append(
            ContractBookDisplayRow(
                row_index=row_index,
                raw_contract_no=raw_contract_text,
                raw_date=raw_date_text,
                contract_no=contract_no,
                year=year,
                ordinal=ordinal,
                contract_date=contract_date,
            )
        )

    rows_by_contract_no: dict[str, list[ContractBookDisplayRow]] = {}
    for row in parsed_rows:
        rows_by_contract_no.setdefault(row.contract_no, []).append(row)

    duplicate_keys = {contract_no for contract_no, rows in rows_by_contract_no.items() if len(rows) > 1}
    clean_rows = [row for row in parsed_rows if row.contract_no not in duplicate_keys]
    for contract_no in sorted(duplicate_keys):
        for row in rows_by_contract_no[contract_no]:
            issues.append(
                _make_issue(
                    row,
                    ContractBookIssueKind.DUPLICATE,
                    "So cong chung bi trung sau khi chuan hoa.",
                )
            )

    clean_rows = _remove_date_sequence_issues(clean_rows, issues)
    clean_rows = _remove_same_day_jump_issues(clean_rows, issues)
    clean_rows = sorted(clean_rows, key=lambda row: (row.year, row.ordinal, row.row_index))
    missing_numbers = _filter_missing_numbers(_build_missing_numbers(clean_rows), issues)
    ordinals = [row.ordinal for row in clean_rows]
    duplicate_count = sum(1 for issue in issues if issue.kind == ContractBookIssueKind.DUPLICATE)

    return ContractBookAnalysis(
        source_path=path,
        display_rows=clean_rows,
        missing_numbers=missing_numbers,
        issue_rows=sorted(issues, key=lambda issue: issue.row_index),
        summary=ContractBookSummary(
            excel_total=excel_total,
            valid_count=len(clean_rows),
            missing_count=len(missing_numbers),
            issue_count=len(issues),
            duplicate_count=duplicate_count,
        ),
        min_ordinal=min(ordinals) if ordinals else None,
        max_ordinal=max(ordinals) if ordinals else None,
    )
