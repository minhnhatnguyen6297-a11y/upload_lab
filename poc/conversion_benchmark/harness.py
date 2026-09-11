from __future__ import annotations

import hashlib
import importlib.metadata
import json
from pathlib import Path
import subprocess
from time import perf_counter
import tracemalloc
from typing import Any, Callable

import fitz
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image

from .router import convert_path


def _baseline_text(path: Path) -> str | None:
    """Read with the existing local stack only where an honest baseline exists."""
    if path.suffix.lower() == ".docx":
        from extract_contract import read_docx

        return read_docx(path)
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            return "\n".join(
                str(value)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows(values_only=True)
                for value in row
                if value is not None
            )
        finally:
            workbook.close()
    return None


def materialize_golden(directory: Path) -> list[dict[str, Any]]:
    """Create only synthetic inputs; callers keep the manifest/report outside production paths."""
    directory.mkdir(parents=True, exist_ok=True)
    pdf_text = directory / "GD-01-text.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Synthetic PDF text")
    pdf_text.write_bytes(pdf.tobytes())
    pdf.close()
    pdf_scan = directory / "GD-02-scan.pdf"
    pdf = fitz.open()
    pdf.new_page()
    pdf_scan.write_bytes(pdf.tobytes())
    pdf.close()
    docx = directory / "GD-03.docx"
    document = Document()
    document.add_paragraph("Synthetic DOCX paragraph")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Value"
    document.save(docx)
    xlsx = directory / "GD-04.xlsx"
    workbook = Workbook()
    workbook.active.title = "Synthetic"
    workbook.active.append(["Field", "Value"])
    workbook.active.append(["Party", "Synthetic A"])
    workbook.save(xlsx)
    legacy = directory / "GD-05.doc"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0synthetic")
    image = directory / "GD-06.png"
    Image.new("RGB", (2, 2), "white").save(image, format="PNG")
    broken = directory / "GD-07-broken.pdf"
    broken.write_bytes(b"%PDF-not-valid")
    template = json.loads(Path(__file__).with_name("golden_manifest.json").read_text(encoding="utf-8"))
    return [{**entry, "path": str(directory / entry.pop("filename"))} for entry in template["sources"]]


def load_persistent_golden(directory: Path | None = None) -> list[dict[str, Any]]:
    """Load committed synthetic fixtures and fail closed on a hash mismatch."""
    golden_dir = Path(__file__).with_name("golden") if directory is None else directory
    manifest_path = Path(__file__).with_name("golden_manifest.json")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    entries: list[dict[str, Any]] = []
    for source in manifest["sources"]:
        filename = source["filename"]
        expected_hash = source.get("sha256")
        if not expected_hash:
            raise ValueError(f"golden manifest entry has no sha256: {filename}")
        path = golden_dir / filename
        if not path.is_file():
            raise FileNotFoundError(f"golden fixture is missing: {path}")
        actual_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        if actual_hash != expected_hash:
            raise ValueError(f"golden fixture hash mismatch: {filename}")
        entries.append({**source, "path": str(path)})
    return entries


def _revision() -> str:
    try:
        return subprocess.run(["git", "rev-parse", "HEAD"], check=True, capture_output=True, text=True).stdout.strip()
    except (OSError, subprocess.SubprocessError):
        return "unknown"


def _versions() -> dict[str, str]:
    versions: dict[str, str] = {}
    for package in ("markitdown", "PyMuPDF", "python-docx", "openpyxl"):
        try:
            versions[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            versions[package] = "not-installed"
    return versions


def _provenance_status(envelope: dict[str, Any]) -> str:
    if any(segment.get("source_ref") is not None for segment in envelope["segments"]):
        return "source-ref"
    if any(warning in envelope["warnings"] for warning in ("provenance_not_asserted_by_markitdown", "legacy_doc_requires_upload_lab_ifilter")):
        return "warning"
    if envelope["route"] == "unsupported":
        return "source-hash"
    if envelope["route"] == "ocr_candidate":
        return "ocr-call-or-warning"
    return "none"


def _recommendation(route: str) -> str:
    if route in {"local"}:
        return "iterate: provenance/source-segment gate remains unmet"
    if route == "legacy_doc_external":
        return "not_adopt: preserve Windows IFilter boundary"
    if route == "ocr_candidate":
        return "not_adopt: OCR needs a separately approved gate"
    return "not_adopt: unsupported input remains structured error"


def run_benchmark(
    manifest: list[dict[str, Any]],
    output_path: Path,
    *,
    converter: Callable[[Path], str] | None = None,
) -> dict[str, Any]:
    results: list[dict[str, Any]] = []
    tracemalloc.start()
    started = perf_counter()
    try:
        for item in manifest:
            sample_started = perf_counter()
            path = Path(item["path"])
            envelope = convert_path(path) if converter is None else convert_path(path, converter=converter)
            text = envelope["content"]["value"]
            baseline_text = _baseline_text(path)
            route_match = envelope["route"] == item["expected_route"]
            text_match = item["expected_text_contains"] in text
            actual_provenance = _provenance_status(envelope)
            result = {
                "sample_id": item["sample_id"],
                "sha256": envelope["source"]["sha256"],
                "route": envelope["route"],
                "route_match": route_match,
                "text_match": text_match,
                "baseline_text_contains": None if baseline_text is None else item["expected_text_contains"] in baseline_text,
                "baseline_available": baseline_text is not None,
                "expected_provenance": item["expected_provenance"],
                "actual_provenance": actual_provenance,
                "provenance_match": actual_provenance == item["expected_provenance"],
                "warnings": envelope["warnings"],
                "errors": envelope["errors"],
                "cloud_call_count": len(envelope["ocr_calls"]),
                "duration_ms": round((perf_counter() - sample_started) * 1000),
                "recommendation": _recommendation(envelope["route"]),
            }
            # An expected unsupported input is a completed measurement when it
            # produces the correct structured route/error boundary.
            result["status"] = "completed" if route_match and text_match and result["provenance_match"] else "partial"
            results.append(result)
    finally:
        _, peak_memory_bytes = tracemalloc.get_traced_memory()
        tracemalloc.stop()
    report = {
        "contract_version": "v0.experimental",
        "poc_revision": _revision(),
        "packages": _versions(),
        "measurement": {"duration_ms": round((perf_counter() - started) * 1000), "peak_python_allocations": peak_memory_bytes},
        "summary": {"total": len(results), "completed": sum(item["status"] == "completed" for item in results), "partial": sum(item["status"] == "partial" for item in results)},
        "cloud_call_count": sum(item["cloud_call_count"] for item in results),
        "estimated_cloud_cost_usd": None,
        "results": results,
        "decision": "review_required",
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report
